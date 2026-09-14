from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path("/home/rogan/Documents/budi/da-project")
SOURCE = ROOT / "output" / "docx" / "Revisi_BAB_II_dan_BAB_III_Homcuts.docx"
OUTPUT = ROOT / "output" / "docx" / "Skripsi_Homcuts_BAB_I_sampai_BAB_V.docx"

sys.path.insert(0, str(ROOT / "tmp"))
import build_revised_bab23 as base  # noqa: E402


def add_centered_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.keep_with_next = True
    base.style_run(p.add_run(title), bold=True)


def add_formula(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    base.style_run(p.add_run(text), italic=True)


def add_bab_i(target: Document) -> None:
    base.add_chapter_heading(target, "BAB I", "PENDAHULUAN")

    base.add_section_heading(target, "1.1", "Latar Belakang")
    base.add_rich_paragraph(
        target,
        "Teknologi informasi berbasis web memungkinkan data operasional dikumpulkan, diolah, dan disajikan melalui satu sistem yang dapat diakses menggunakan peramban. Bagi usaha jasa, manfaat sistem tidak hanya berupa penyajian informasi kepada pelanggan, tetapi juga integrasi kegiatan pelayanan dan pencatatan manajemen. Sistem informasi manajemen yang terintegrasi mengurangi pencatatan berulang dan membantu penyediaan informasi yang konsisten untuk kegiatan operasional serta pengambilan keputusan (Laudon & Laudon, 2022).",
    )
    base.add_rich_paragraph(
        target,
        "Homcuts merupakan usaha barbershop yang melayani pelanggan berdasarkan ketersediaan barber dan waktu pelayanan. Kegiatan utamanya mencakup penyampaian informasi usaha, pemesanan layanan, penerimaan pelanggan walk-in, penjualan produk perawatan rambut, penerimaan pembayaran tunai, dan pencatatan transaksi. Setiap layanan memiliki durasi, sementara setiap barber memiliki jam kerja. Kondisi tersebut menyebabkan pengelolaan jadwal tidak cukup dilakukan dengan mencatat jam mulai saja, tetapi harus memperhitungkan keseluruhan interval layanan.",
        italic_terms=["barbershop", "barber", "walk-in"],
    )
    base.add_rich_paragraph(
        target,
        "Apabila booking pelanggan dan layanan walk-in dicatat melalui mekanisme yang terpisah, dua pelanggan dapat memperoleh nomor antrean yang sama atau menggunakan barber pada waktu yang saling bertabrakan. Perubahan jadwal oleh admin juga dapat menimbulkan perbedaan informasi apabila halaman pelanggan tidak membaca sumber data yang sama. Penelitian tentang aplikasi booking barbershop menunjukkan bahwa pemesanan berbasis web membantu pengelolaan reservasi, sedangkan penerapan POS berbasis web membantu pemusatan transaksi di titik pelayanan (Firmansyah dkk., 2023; Ngatini & Cahyanti, 2024). Namun, kebutuhan Homcuts tidak hanya mencakup booking atau POS secara terpisah, melainkan penyatuan keduanya dalam jadwal dan antrean pelayanan yang sama.",
        italic_terms=["booking", "walk-in", "barber", "barbershop"],
    )
    base.add_rich_paragraph(
        target,
        "Permasalahan lain terdapat pada pencatatan penjualan produk dan pembayaran. Pelanggan dapat memesan produk melalui situs untuk diambil di lokasi, sedangkan kasir juga dapat mencatat penjualan langsung. Karena pembayaran pada ruang lingkup penelitian dilakukan secara tunai, transaksi harus tetap berstatus belum dibayar sampai uang diterima dan dikonfirmasi oleh kasir. Status pembayaran perlu dibedakan dari status pemenuhan layanan atau pesanan agar transaksi yang sudah lunas tetapi belum selesai tetap dapat dipantau. Catatan transaksi tersebut juga harus menjadi dasar ringkasan penerimaan usaha dan riwayat aktivitas.",
    )
    base.add_rich_paragraph(
        target,
        "Berdasarkan kebutuhan tersebut, penelitian ini merancang dan membangun Sistem Informasi Manajemen Homcuts berbasis web. Sistem menyediakan halaman publik untuk informasi profil, barber, layanan, produk, galeri, kontak, booking, dan pemantauan status. Panel admin menyediakan autentikasi, ringkasan operasional, POS, pengelolaan booking dan transaksi, konfirmasi pembayaran tunai, notifikasi, serta CRUD data utama. Booking dan POS menggunakan pemeriksaan interval waktu serta sumber nomor antrean yang sama, sedangkan perubahan status dibaca secara berkala oleh halaman pelanggan dan admin sehingga informasi dapat diperbarui tanpa muat ulang manual.",
        italic_terms=["barber", "booking", "CRUD"],
    )
    base.add_rich_paragraph(
        target,
        "Sistem dikembangkan menggunakan Laravel dengan pola Model-View-Controller, Blade, Tailwind CSS, JavaScript, Vite, MySQL, dan Docker Compose. Pengujian fungsional dilakukan dengan pendekatan black-box terhadap halaman publik, autentikasi, pengelolaan data, booking, POS, antrean, jadwal, transaksi produk, pembayaran tunai, notifikasi, dan pembaruan status. Rancangan evaluasi kegunaan dan efisiensi juga disiapkan untuk pelaksanaan bersama pengguna pada tahap evaluasi lapangan.",
        italic_terms=["Model-View-Controller", "Blade", "JavaScript", "black-box", "booking"],
    )

    base.add_section_heading(target, "1.2", "Rumusan Masalah")
    base.add_rich_paragraph(target, "Berdasarkan latar belakang tersebut, rumusan masalah penelitian ini adalah:")
    base.add_list_item(target, "1)", "Bagaimana merancang dan membangun sistem informasi manajemen Homcuts berbasis web yang mengintegrasikan informasi publik, booking, pemesanan produk, POS, pembayaran tunai, pengelolaan data, dan pelaporan transaksi?")
    base.add_list_item(target, "2)", "Bagaimana menerapkan jadwal dan nomor antrean terpadu untuk booking serta pelanggan walk-in dengan mempertimbangkan durasi layanan, jam operasional, jam kerja barber, dan benturan interval waktu?")
    base.add_list_item(target, "3)", "Bagaimana menerapkan konfirmasi pembayaran tunai dan sinkronisasi perubahan status pada halaman pelanggan serta admin tanpa muat ulang manual?")
    base.add_list_item(target, "4)", "Bagaimana hasil pengujian fungsional sistem informasi manajemen Homcuts menggunakan metode black-box?")

    base.add_section_heading(target, "1.3", "Batasan Masalah")
    base.add_rich_paragraph(target, "Batasan yang digunakan agar penelitian tetap terarah adalah sebagai berikut:")
    limits = [
        "Objek penelitian adalah Homcuts Barbershop dan sistem dibangun dalam bentuk aplikasi web.",
        "Aktor sistem terdiri atas pelanggan dan Admin/Kasir. Pelanggan tidak diwajibkan membuat akun, sedangkan akses admin menggunakan autentikasi berbasis sesi.",
        "Halaman pelanggan mencakup beranda, booking, toko produk, galeri, tentang kami, kontak, dan status transaksi.",
        "Panel admin mencakup ringkasan, POS, booking, transaksi, notifikasi, akun, serta CRUD data barber, layanan, produk, galeri, pesan, dan pengaturan situs.",
        "Booking dan layanan walk-in menggunakan jadwal serta nomor antrean terpadu. Ketersediaan dihitung dari durasi layanan, jam operasional pukul 07.00–22.00, jam kerja barber, dan data layanan aktif lain.",
        "Pembayaran yang dibahas hanya pembayaran tunai manual. Sistem tidak mengaktifkan QRIS, payment gateway, pengiriman produk, penggajian, perpajakan, atau akuntansi buku besar.",
        "Produk yang dipesan melalui situs diambil langsung di Homcuts. Sistem mengelola keranjang, checkout, stok, konfirmasi pembayaran, dan status pengambilan.",
        "Pengujian utama pada naskah ini adalah pengujian fungsional black-box di lingkungan Docker dengan basis data MySQL khusus pengujian. Instrumen SUS dan efisiensi waktu disiapkan untuk evaluasi lapangan, tetapi skor tidak dinyatakan tanpa data responden dan pengukuran nyata.",
    ]
    for i, item in enumerate(limits, 1):
        base.add_list_item(target, f"{i})", item)

    base.add_section_heading(target, "1.4", "Tujuan Penelitian")
    base.add_rich_paragraph(target, "Tujuan penelitian yang disusun sesuai rumusan masalah adalah:")
    objectives = [
        "Merancang dan membangun sistem informasi manajemen Homcuts berbasis web yang mengintegrasikan informasi publik, booking, pemesanan produk, POS, pembayaran tunai, pengelolaan data, dan pelaporan transaksi.",
        "Menerapkan jadwal dan nomor antrean terpadu untuk booking serta pelanggan walk-in dengan validasi durasi, jam operasional, jam kerja barber, dan benturan interval waktu.",
        "Menerapkan konfirmasi pembayaran tunai serta sinkronisasi perubahan status pada halaman pelanggan dan admin tanpa muat ulang manual.",
        "Mengetahui kesesuaian fungsi sistem terhadap kebutuhan melalui pengujian black-box.",
    ]
    for i, item in enumerate(objectives, 1):
        base.add_list_item(target, f"{i})", item)

    base.add_section_heading(target, "1.5", "Manfaat Penelitian")
    base.add_section_heading(target, "1.5.1", "Bagi Homcuts", level=2)
    base.add_list_item(target, "1)", "Menyediakan pencatatan jadwal, antrean, transaksi, pembayaran, dan data utama dalam satu sistem.")
    base.add_list_item(target, "2)", "Membantu kasir melayani pelanggan booking maupun walk-in serta memantau pekerjaan yang belum selesai.")
    base.add_list_item(target, "3)", "Menyediakan riwayat transaksi dan ringkasan operasional yang dapat digunakan pemilik dalam memantau usaha.")
    base.add_section_heading(target, "1.5.2", "Bagi Pelanggan", level=2)
    base.add_list_item(target, "1)", "Memudahkan pelanggan memperoleh informasi layanan, barber, produk, galeri, kontak, dan lokasi Homcuts.")
    base.add_list_item(target, "2)", "Memungkinkan pelanggan memilih waktu layanan yang tersedia dan memantau perubahan jadwal maupun status pembayaran.")
    base.add_list_item(target, "3)", "Memudahkan pemesanan produk untuk diambil dan dibayar langsung di lokasi.")
    base.add_section_heading(target, "1.5.3", "Bagi Peneliti dan Pengembang", level=2)
    base.add_rich_paragraph(target, "Penelitian ini dapat menjadi contoh penerapan Laravel, basis data relasional, transaksi tunai, jadwal interval, antrean terpadu, dan pengujian fungsional pada sistem informasi manajemen usaha jasa skala kecil.")


def page_break_only(element) -> bool:
    if element.tag != qn("w:p"):
        return False
    breaks = list(element.iter(qn("w:br")))
    return bool(breaks) and all(br.get(qn("w:type")) in (None, "page") for br in breaks)


def remove_old_bibliography(doc: Document) -> None:
    body = doc._element.body
    marker = None
    for paragraph in doc.paragraphs:
        if "DAFTAR RUJUKAN" in paragraph.text:
            marker = paragraph._element
            break
    if marker is None:
        return
    previous = marker.getprevious()
    children = list(body)
    start = children.index(marker)
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    if previous is not None and previous.getparent() is body and page_break_only(previous):
        body.remove(previous)


def prepend_bab_i(doc: Document) -> None:
    first = Document()
    base.configure_document(first)
    base.set_page_number_start(first.sections[0], 1)
    add_bab_i(first)
    first.add_page_break()
    body = doc._element.body
    insert_at = 0
    for element in list(first._element.body):
        if element.tag == qn("w:sectPr"):
            continue
        body.insert(insert_at, deepcopy(element))
        insert_at += 1


def normalize_chapter_breaks(doc: Document) -> None:
    starts = ("BAB II", "BAB III", "BAB IV", "BAB V", "DAFTAR PUSTAKA")
    for paragraph in list(doc.paragraphs):
        compact = " ".join(paragraph.text.split())
        if not compact.startswith(starts):
            continue
        paragraph.paragraph_format.page_break_before = True
        previous = paragraph._element.getprevious()
        if previous is not None and page_break_only(previous):
            previous.getparent().remove(previous)


def add_bab_iv(doc: Document) -> None:
    doc.add_page_break()
    base.add_chapter_heading(doc, "BAB IV", "HASIL DAN PEMBAHASAN")
    base.add_rich_paragraph(
        doc,
        "Bab ini menyajikan hasil implementasi dan pengujian Sistem Informasi Manajemen Homcuts. Hasil disusun berdasarkan kebutuhan dan rancangan pada Bab III. Pembahasan kemudian menghubungkan temuan tersebut dengan rumusan masalah serta konsep pada Bab II.",
    )

    base.add_section_heading(doc, "4.1", "Hasil Implementasi Sistem")
    base.add_section_heading(doc, "4.1.1", "Lingkungan Implementasi", level=2)
    base.add_rich_paragraph(
        doc,
        "Sistem diimplementasikan sebagai aplikasi web berarsitektur MVC. Laravel mengelola route, autentikasi, validasi, logika bisnis, akses basis data, dan penyimpanan berkas. Blade, Tailwind CSS, JavaScript, dan Vite membentuk antarmuka pelanggan serta admin. Aplikasi, MySQL, dan scheduler dijalankan sebagai layanan terpisah melalui Docker Compose. Rincian lingkungan yang digunakan pada verifikasi akhir disajikan pada Tabel 4.1.",
        italic_terms=["route", "Blade", "JavaScript", "scheduler"],
    )
    base.add_table(
        doc,
        "Tabel 4.1 Lingkungan implementasi dan pengujian",
        ["Komponen", "Versi/konfigurasi", "Peran"],
        [
            ["PHP", "8.4.24", "Menjalankan aplikasi sisi server."],
            ["Laravel", "13.26.1", "Kerangka MVC, autentikasi, validasi, ORM, dan pengujian."],
            ["MySQL", "8.4", "Menyimpan data utama dan transaksi."],
            ["Tailwind CSS", "4.3.3", "Menyusun tampilan responsif."],
            ["Vite", "8.2.2", "Membangun aset CSS dan JavaScript."],
            ["PHPUnit", "12.5.33", "Menjalankan pengujian otomatis."],
            ["Docker Compose", "web, db, scheduler", "Menyeragamkan lingkungan aplikasi dan layanan pendukung."],
        ],
        widths=[3.2, 4.0, 7.2],
        font_size=8.6,
    )

    base.add_section_heading(doc, "4.1.2", "Modul Pelanggan", level=2)
    base.add_rich_paragraph(
        doc,
        "Antarmuka pelanggan terdiri atas halaman Beranda, Booking, Toko Produk, Galeri, Tentang Kami, Kontak, dan Status Transaksi. Data barber, layanan, produk, dan galeri dibaca dari basis data sehingga perubahan yang dilakukan admin dapat ditampilkan pada halaman publik. Pelanggan dapat melihat foto serta informasi layanan, memilih produk ke dalam keranjang, mengirim pesan, dan membuat booking tanpa akun.",
        italic_terms=["barber", "Booking", "booking"],
    )
    base.add_rich_paragraph(
        doc,
        "Formulir booking menerima tanggal dan waktu spesifik. Sistem menghitung waktu selesai berdasarkan durasi layanan, memeriksa batas operasional pukul 07.00–22.00, memeriksa jam kerja barber, dan menolak interval yang bertabrakan. Setelah booking tersimpan, sistem membuat transaksi, pembayaran tunai berstatus belum dibayar, nomor antrean, serta notifikasi admin. Halaman status pelanggan menampilkan kode transaksi, jadwal terbaru, barber, nomor antrean, nominal, status pembayaran, dan status pelayanan.",
        italic_terms=["booking", "barber"],
    )

    base.add_section_heading(doc, "4.1.3", "Modul Admin dan Kasir", level=2)
    base.add_rich_paragraph(
        doc,
        "Panel admin dilindungi login dan pemeriksaan hak akses. Navigasi operasional menempatkan Ringkasan, Kasir POS, Booking, dan Transaksi pada bagian awal. Dashboard menyajikan pekerjaan tertunda dan ringkasan penerimaan, sedangkan POS difokuskan pada pencatatan transaksi. Tabel admin dapat diurutkan langsung melalui judul kolom dan bagian data operasional melakukan pembaruan berkala tanpa muat ulang halaman secara manual.",
        italic_terms=["Dashboard"],
    )
    base.add_rich_paragraph(
        doc,
        "Melalui POS, kasir memasukkan pelanggan walk-in, waktu pelayanan, barber, layanan, dan produk tambahan. Jadwal POS diperiksa terhadap booking dan transaksi layanan lain menggunakan aturan interval yang sama. Nomor antrean layanan dihasilkan oleh QueueNumberService untuk seluruh saluran. Produk yang tidak memuat layanan tidak memperoleh nomor antrean karena tidak menggunakan sumber daya barber.",
        italic_terms=["walk-in", "barber", "booking", "QueueNumberService"],
    )
    base.add_rich_paragraph(
        doc,
        "Pembayaran booking dapat dikonfirmasi pada halaman Booking atau Transaksi, sedangkan pembayaran pesanan produk dikonfirmasi pada halaman Transaksi. Setelah kasir menyatakan uang diterima, PaymentService mengisi waktu dan petugas konfirmasi, mengubah status pembayaran menjadi lunas, serta menyelaraskan status pelayanan. Admin juga dapat mengelola data barber, layanan, produk, galeri, pesan, dan pengaturan situs, termasuk unggah foto.",
        italic_terms=["booking", "PaymentService", "barber"],
    )

    base.add_section_heading(doc, "4.1.4", "Keterkaitan Modul dan Data", level=2)
    base.add_table(
        doc,
        "Tabel 4.2 Hasil implementasi modul sistem",
        ["Modul", "Data utama", "Hasil implementasi"],
        [
            ["Informasi publik", "barbers, services, products, gallery_entries, site_settings", "Data aktif ditampilkan pada halaman pelanggan dan dapat dikelola admin."],
            ["Booking", "bookings, orders, payments", "Jadwal, transaksi, pembayaran tertunda, dan status pelanggan terbentuk dalam satu alur."],
            ["POS/walk-in", "orders, order_items, payments", "Kasir mencatat layanan/produk, jadwal, antrean, dan pembayaran tunai."],
            ["Antrean terpadu", "queue_sequences, orders", "Booking dan POS layanan memperoleh nomor unik dari sumber yang sama."],
            ["Penjualan produk", "products, orders, order_items, payments", "Total dihitung server, stok diperiksa, dan pesanan disiapkan untuk diambil."],
            ["Administrasi", "users, notifications, contact_messages", "Akses dilindungi, aktivitas penting diberitahukan, dan pesan pelanggan tersimpan."],
            ["Pelaporan", "orders, payments", "Riwayat transaksi menyatukan booking, POS, dan pesanan produk."],
        ],
        widths=[3.0, 5.0, 6.4],
        font_size=8.2,
    )

    base.add_section_heading(doc, "4.2", "Hasil Pengujian Black-Box")
    base.add_section_heading(doc, "4.2.1", "Pelaksanaan Pengujian", level=2)
    base.add_rich_paragraph(
        doc,
        "Pengujian dijalankan pada container aplikasi sementara dengan APP_ENV bernilai testing dan basis data MySQL barbershop_testing yang terpisah dari data utama. Cache dan sesi menggunakan penyimpanan sementara. Setiap pengujian menyiapkan kondisi awal, mengirim masukan melalui antarmuka HTTP atau layanan aplikasi, kemudian memeriksa respons, perubahan basis data, dan larangan yang berlaku. Cara ini mencegah hasil pengujian mengubah data operasional Homcuts.",
        italic_terms=["container", "testing", "barbershop_testing"],
    )
    base.add_rich_paragraph(
        doc,
        "Pengujian akhir pada 13 September 2026 menggunakan PHPUnit 12.5.33 menghasilkan 48 metode uji dan 595 assertion. Seluruh metode uji lulus dalam waktu 27,106 detik dengan penggunaan memori 62,50 MB. Tabel 4.3 merangkum hasil berdasarkan kelompok fungsi.",
        italic_terms=["assertion"],
    )
    base.add_table(
        doc,
        "Tabel 4.3 Ringkasan hasil pengujian otomatis",
        ["Kelompok", "Jumlah uji", "Hasil", "Cakupan utama"],
        [
            ["Halaman publik dan alur pelanggan", "9", "9 lulus", "Halaman publik, merek, galeri/layanan, modal, booking, jam toko, total server, dan kontak."],
            ["Admin, autentikasi, CRUD, dan antarmuka", "13", "13 lulus", "Hak akses, login, dashboard, pembaruan langsung, navigasi, sorting, CRUD, dan unggah foto."],
            ["Booking, pembayaran, notifikasi, dan integritas", "18", "18 lulus", "Antrean, reschedule, overlap, expiry, idempotensi, jam kerja, nilai pembayaran, dan notifikasi."],
            ["POS dan integrasi layanan", "7", "7 lulus", "POS tunai, antrean layanan, jadwal bersama, panel tertunda, dan pencegahan transaksi ganda."],
            ["Uji dasar", "1", "1 lulus", "Pemeriksaan dasar lingkungan unit test."],
            ["Total", "48", "48 lulus", "595 assertion; tingkat kelulusan 100%."],
        ],
        widths=[4.1, 2.0, 2.2, 6.1],
        font_size=8.0,
    )
    add_formula(doc, "Tingkat kelulusan = (48 / 48) × 100% = 100%")

    base.add_section_heading(doc, "4.2.2", "Hasil Kasus Uji Fungsional", level=2)
    base.add_rich_paragraph(doc, "Hasil terhadap skenario inti pada Tabel 3.5 disajikan pada Tabel 4.4. Kolom hasil aktual diringkas dari respons aplikasi dan pemeriksaan basis data pada pengujian otomatis.")
    base.add_table(
        doc,
        "Tabel 4.4 Hasil pengujian black-box fungsi inti",
        ["Kode", "Hasil aktual", "Status"],
        [
            ["B-01", "Seluruh halaman publik merespons berhasil dan menampilkan data aktif.", "Lulus"],
            ["B-02", "Booking, order, pembayaran tertunda, notifikasi, dan nomor antrean tersimpan.", "Lulus"],
            ["B-03", "Waktu di luar jam kerja dan interval bertabrakan ditolak dengan pesan validasi.", "Lulus"],
            ["B-04", "Slot yang dimulai tepat setelah layanan sebelumnya berakhir dapat disimpan.", "Lulus"],
            ["B-05", "POS yang bertabrakan dengan booking ditolak tanpa membentuk transaksi baru.", "Lulus"],
            ["B-06", "Booking dan POS layanan memperoleh nomor antrean global yang berbeda.", "Lulus"],
            ["B-07", "Perubahan jadwal admin tampil pada status pelanggan tanpa mengganti nomor antrean.", "Lulus"],
            ["B-08", "Checkout produk menghitung total pada server, menyimpan item, mengurangi stok, dan menunggu tunai.", "Lulus"],
            ["B-09", "Jumlah melebihi stok ditolak dan persediaan tidak berubah.", "Lulus"],
            ["B-10", "Konfirmasi tunai booking mengubah pembayaran menjadi lunas dan booking terkonfirmasi.", "Lulus"],
            ["B-11", "Konfirmasi tunai produk mengubah pesanan menjadi lunas dan siap diambil.", "Lulus"],
            ["B-12", "Konfirmasi berulang tidak menggandakan penerimaan atau perubahan data.", "Lulus"],
            ["B-13", "Tamu yang membuka panel admin diarahkan ke login; nonadmin ditolak.", "Lulus"],
            ["B-14", "Data login salah ditolak; akun admin dapat masuk dan keluar.", "Lulus"],
            ["B-15", "CRUD produk dan unggah foto barber/galeri tervalidasi dan tersimpan.", "Lulus"],
            ["B-16", "Seluruh kolom tabel yang terlihat dapat diurutkan naik dan turun melalui header.", "Lulus"],
            ["B-17", "Data operasional dan notifikasi diperbarui berkala tanpa refresh manual.", "Lulus"],
        ],
        widths=[1.2, 11.1, 2.1],
        font_size=7.0,
    )

    base.add_section_heading(doc, "4.3", "Hasil Pengujian User Acceptance Testing")
    base.add_section_heading(doc, "4.3.1", "Evaluasi Penerimaan Formatif", level=2)
    base.add_rich_paragraph(
        doc,
        "Selama pengembangan, pengguna memberikan evaluasi terhadap alur yang sedang berjalan. Evaluasi tersebut bersifat formatif karena digunakan untuk memperbaiki kebutuhan dan implementasi sebelum UAT formal. Masukan yang dapat ditelusuri meliputi perubahan jam booking, penyatuan antrean dan jadwal, penyederhanaan pembayaran menjadi tunai, lokasi konfirmasi pembayaran, pembaruan status tanpa muat ulang manual, serta susunan navigasi dan tabel admin. Tabel 4.5 menunjukkan tindak lanjut terhadap masukan tersebut.",
        italic_terms=["booking", "UAT"],
    )
    base.add_table(
        doc,
        "Tabel 4.5 Hasil evaluasi penerimaan formatif",
        ["Temuan pengguna", "Perbaikan yang diterapkan", "Bukti verifikasi"],
        [
            ["Pilihan waktu tidak boleh 24 jam.", "Jam layanan dibatasi pukul 07.00–22.00 dan mengikuti durasi serta shift barber.", "B-03 dan pengujian aturan jam kerja lulus."],
            ["Nomor antrean booking dan POS pernah sama.", "QueueNumberService menggunakan satu sumber nomor untuk seluruh transaksi layanan.", "B-06 dan pengujian keunikan antrean lulus."],
            ["Jadwal POS harus memengaruhi slot booking.", "Booking dan POS memakai pemeriksaan interval serta barber yang sama.", "B-04 dan B-05 lulus."],
            ["QRIS belum diperlukan.", "Alur aktif dibatasi pada pembayaran tunai manual; permintaan QRIS ditolak.", "Pengujian penolakan QRIS lulus."],
            ["Pembayaran produk perlu dikonfirmasi.", "Konfirmasi tunai produk ditempatkan pada halaman Transaksi.", "B-11 dan pengujian status siap diambil lulus."],
            ["Status tidak boleh menunggu refresh manual.", "Halaman status, tabel admin, dan notifikasi mengambil pembaruan berkala.", "B-07 dan B-17 lulus."],
            ["Navigasi dan tabel admin perlu lebih ringkas.", "Urutan menu disesuaikan dan sorting dilakukan langsung dari judul kolom.", "B-16 serta pengujian navigasi admin lulus."],
        ],
        widths=[4.2, 6.0, 4.2],
        font_size=7.8,
    )

    base.add_section_heading(doc, "4.3.2", "Status UAT Formal", level=2)
    base.add_rich_paragraph(
        doc,
        "Sepuluh skenario UAT pada Tabel 3.6 telah disusun untuk mewakili Pelanggan serta Admin/Kasir. Prasyarat teknisnya telah terpenuhi karena fungsi terkait lulus pada pengujian black-box. Akan tetapi, dokumen sumber belum memuat keputusan diterima atau tidak diterima dari peserta, catatan pelaksanaan, tanggal UAT, maupun lembar persetujuan pihak Homcuts. Oleh karena itu, persentase penerimaan dan keputusan akhir UAT belum dapat dihitung. Status setiap komponen pelaksanaan dirangkum pada Tabel 4.6.",
        italic_terms=["UAT", "black-box"],
    )
    base.add_table(
        doc,
        "Tabel 4.6 Status pelaksanaan UAT formal",
        ["Komponen", "Status", "Keterangan"],
        [
            ["Skenario dan kriteria penerimaan", "Siap", "Sepuluh skenario tersedia pada Tabel 3.6."],
            ["Stabilitas fungsi inti", "Siap", "Seluruh 48 pengujian dan 595 assertion lulus."],
            ["Data dan lingkungan uji", "Siap", "Data uji dapat disiapkan terpisah dari data operasional."],
            ["Pelaksanaan oleh Pelanggan", "Belum terdokumentasi", "Belum tersedia keputusan dan catatan peserta pelanggan."],
            ["Pelaksanaan oleh Admin/Kasir", "Belum terdokumentasi", "Belum tersedia keputusan dan catatan pihak Homcuts."],
            ["Persentase penerimaan", "Belum dihitung", "Perhitungan menunggu seluruh keputusan skenario."],
            ["Persetujuan akhir", "Belum tersedia", "Lembar penerimaan harus ditandatangani pihak yang berwenang."],
        ],
        widths=[4.2, 3.4, 6.8],
        font_size=8.1,
    )
    base.add_rich_paragraph(
        doc,
        "Status tersebut berarti sistem telah siap memasuki UAT formal, tetapi belum dapat dinyatakan diterima oleh pengguna hanya berdasarkan pengujian otomatis. Setelah pelaksanaan, keputusan pada setiap skenario perlu dicatat. Apabila seluruh skenario kritis diterima dan persentase keseluruhan mencapai sekurang-kurangnya 80 persen, sistem dapat dinyatakan diterima; jika tidak, temuan diperbaiki dan skenario terkait diuji ulang.",
        italic_terms=["UAT"],
    )

    base.add_section_heading(doc, "4.4", "Pembahasan")
    base.add_section_heading(doc, "4.4.1", "Integrasi Informasi dan Operasional", level=2)
    base.add_rich_paragraph(
        doc,
        "Hasil implementasi menjawab rumusan masalah pertama. Informasi publik dan aktivitas operasional menggunakan sumber data MySQL yang sama. Data layanan, barber, produk, galeri, dan pengaturan yang dikelola admin menjadi masukan bagi halaman pelanggan. Booking, POS, penjualan produk, pembayaran, notifikasi, dan laporan transaksi dihubungkan melalui relasi basis data. Integrasi tersebut sesuai dengan konsep sistem informasi manajemen, yaitu data dari satu proses dapat digunakan kembali untuk proses lain sehingga pencatatan ganda dan perbedaan informasi dapat dikurangi (Laudon & Laudon, 2022).",
        italic_terms=["barber", "Booking"],
    )
    base.add_rich_paragraph(
        doc,
        "Berbeda dari aplikasi yang hanya berfokus pada reservasi, sistem Homcuts menghubungkan booking dengan order sejak booking dibuat. Oleh karena itu, booking yang belum dibayar telah terlihat sebagai pekerjaan tertunda, sedangkan booking yang selesai dapat langsung menjadi bagian riwayat transaksi. Pola ini memperluas manfaat booking berbasis web yang dilaporkan Firmansyah dkk. (2023) dengan menambahkan pengelolaan transaksi dan status pembayaran dalam alur yang sama.",
        italic_terms=["booking", "order"],
    )

    base.add_section_heading(doc, "4.4.2", "Jadwal dan Antrean Terpadu", level=2)
    base.add_rich_paragraph(
        doc,
        "Rumusan masalah kedua dijawab melalui BookingAvailabilityService dan QueueNumberService. Ketersediaan tidak ditentukan hanya dari kesamaan jam mulai, tetapi dari perpotongan interval waktu mulai dan waktu selesai. Hasil B-03, B-04, B-05, dan pengujian jam kerja menunjukkan bahwa sistem menolak interval yang beririsan, menerima slot yang bersebelahan, dan menghormati batas toko serta shift barber. Aturan yang sama dipanggil oleh alur pelanggan, admin, dan POS sehingga jadwal tidak terpisah menurut saluran pencatatan.",
        italic_terms=["BookingAvailabilityService", "QueueNumberService", "barber"],
    )
    base.add_rich_paragraph(
        doc,
        "Nomor antrean disimpan pada transaksi layanan dan dihasilkan dari satu urutan. Hasil B-06 membuktikan bahwa booking serta POS memperoleh nomor yang berbeda, sedangkan transaksi produk tanpa layanan tidak mengambil nomor antrean. Dengan demikian, identitas antrean merepresentasikan urutan layanan fisik, bukan urutan seluruh transaksi penjualan. Perubahan jadwal mempertahankan nomor antrean yang sama agar identitas pelanggan tidak berubah hanya karena penyesuaian waktu.",
        italic_terms=["booking"],
    )

    base.add_section_heading(doc, "4.4.3", "Pembayaran Tunai dan Sinkronisasi Status", level=2)
    base.add_rich_paragraph(
        doc,
        "Rumusan masalah ketiga dijawab dengan pemisahan status transaksi dan status pembayaran. Ketika booking atau pesanan dibuat, pembayaran belum dianggap diterima. Konfirmasi hanya dilakukan oleh Admin/Kasir setelah uang tunai diterima. Hasil B-10 dan B-11 menunjukkan bahwa konfirmasi memutakhirkan pembayaran sekaligus tahap pelayanan yang sesuai. Pengujian idempotensi memastikan tindakan yang sama tidak mencatat penerimaan dua kali.",
        italic_terms=["booking"],
    )
    base.add_rich_paragraph(
        doc,
        "Pembaruan tanpa muat ulang manual diterapkan melalui permintaan berkala dari JavaScript ke endpoint status. Server tetap menjadi sumber kebenaran; halaman hanya mengganti bagian tampilan ketika data terbaru berbeda. Pendekatan ini lebih sederhana untuk lingkungan penelitian dibanding infrastruktur komunikasi waktu nyata penuh, tetapi telah memenuhi kebutuhan agar perubahan pembayaran, booking, tabel admin, dan notifikasi dapat terlihat tanpa pengguna menekan tombol refresh.",
        italic_terms=["JavaScript", "endpoint", "booking"],
    )

    base.add_section_heading(doc, "4.4.4", "Kesesuaian Fungsional dan Kesiapan Penerimaan", level=2)
    base.add_rich_paragraph(
        doc,
        "Rumusan masalah keempat dijawab oleh hasil black-box. Tingkat kelulusan 100% menunjukkan bahwa keluaran pada kondisi yang diuji sesuai dengan kebutuhan yang dirumuskan. Pengujian tidak hanya memeriksa halaman berhasil dibuka, tetapi juga memeriksa perubahan basis data, pembatasan hak akses, perhitungan total pada server, stok, benturan interval, expiry, idempotensi, serta konsistensi transaksi. Hasil tersebut memberikan bukti bahwa fungsi inti dapat dijalankan pada lingkungan pengujian yang ditetapkan.",
        italic_terms=["black-box", "expiry"],
    )
    base.add_rich_paragraph(
        doc,
        "Meskipun seluruh pengujian lulus, angka 100% tidak berarti sistem bebas kesalahan pada seluruh kondisi penggunaan. Kesimpulan hanya berlaku untuk 48 metode uji, 595 assertion, versi kode, konfigurasi, dan lingkungan yang digunakan. Pengujian di perangkat lain, beban banyak pengguna, gangguan jaringan, serta observasi penggunaan nyata tetap diperlukan sebelum penerapan produksi secara luas.",
        italic_terms=["assertion"],
    )
    base.add_rich_paragraph(
        doc,
        "Evaluasi formatif menunjukkan bahwa masukan pengguna telah menghasilkan perubahan konkret pada sistem dan perubahan tersebut lulus pengujian fungsional. Namun, verifikasi teknis bukan pengganti keputusan penerimaan. UAT formal tetap diperlukan untuk membuktikan bahwa pelanggan serta Admin/Kasir dapat menyelesaikan tugas dan menerima hasilnya dalam konteks penggunaan Homcuts yang sebenarnya.",
        italic_terms=["UAT"],
    )

    base.add_section_heading(doc, "4.5", "Keterbatasan Pengujian")
    base.add_rich_paragraph(
        doc,
        "Bab III telah menyiapkan instrumen UAT, SUS, dan perbandingan efisiensi waktu. Namun, dokumen sumber belum memuat lembar keputusan UAT, jawaban sepuluh butir SUS, maupun catatan waktu proses manual dan proses sistem. Oleh karena itu, penelitian ini tidak menyatakan penerimaan akhir, skor kegunaan, atau persentase penghematan waktu yang belum diukur. Instrumen tersebut dapat langsung digunakan pada evaluasi lapangan; hasilnya perlu ditambahkan pada Bab IV setelah data nyata diperoleh. Pembatasan ini menjaga agar pembahasan dan kesimpulan tidak dibangun dari keputusan atau angka yang direkayasa.",
        italic_terms=["UAT", "SUS"],
    )


def add_bab_v(doc: Document) -> None:
    doc.add_page_break()
    base.add_chapter_heading(doc, "BAB V", "PENUTUP")
    base.add_section_heading(doc, "5.1", "Kesimpulan")
    base.add_rich_paragraph(doc, "Berdasarkan hasil implementasi, pengujian, dan pembahasan, kesimpulan penelitian ini adalah:")
    conclusions = [
        "Sistem Informasi Manajemen Homcuts berbasis web telah dibangun dengan mengintegrasikan informasi publik, booking, pemesanan produk, POS, pembayaran tunai, pengelolaan data utama, notifikasi, dan riwayat transaksi. Seluruh modul menggunakan data terpusat pada MySQL dan dapat dijalankan melalui lingkungan Docker Compose.",
        "Jadwal booking dan pelanggan walk-in telah disatukan. Sistem menghitung interval berdasarkan durasi layanan, menerapkan jam operasional serta jam kerja barber, menolak jadwal bertabrakan, dan menghasilkan nomor antrean layanan yang unik dari sumber yang sama.",
        "Konfirmasi pembayaran tunai dapat dilakukan oleh Admin/Kasir sesuai jenis transaksi. Pembayaran booking dapat dikonfirmasi dari halaman Booking atau Transaksi, sedangkan pembayaran produk dikonfirmasi dari halaman Transaksi. Perubahan jadwal dan status dapat ditampilkan pada halaman pelanggan maupun admin tanpa muat ulang manual melalui pembaruan data berkala.",
        "Pengujian fungsional black-box pada basis data MySQL khusus pengujian menghasilkan 48 pengujian dan 595 assertion yang seluruhnya lulus. Hasil tersebut menunjukkan bahwa fungsi yang diuji telah sesuai dengan keluaran yang diharapkan pada lingkungan pengujian yang digunakan.",
        "Evaluasi penerimaan formatif menghasilkan perbaikan pada jam layanan, antrean dan jadwal terpadu, pembayaran tunai, pembaruan status, serta antarmuka admin. Sepuluh skenario UAT formal telah disiapkan, tetapi keputusan penerimaan akhir belum dapat dinyatakan sebelum skenario dijalankan dan disetujui oleh Pelanggan serta pihak Homcuts.",
    ]
    for i, item in enumerate(conclusions, 1):
        base.add_list_item(doc, f"{i})", item)

    base.add_section_heading(doc, "5.2", "Saran")
    suggestions = [
        "Sebelum naskah akhir dinyatakan lengkap, UAT formal, SUS, dan pengujian efisiensi waktu perlu dilaksanakan bersama pelanggan serta Admin/Kasir. Keputusan skenario, lembar persetujuan, jawaban responden, dan catatan waktu nyata kemudian diolah menggunakan prosedur pada Bab III dan ditambahkan ke Bab IV.",
        "Pada penerapan produksi, Homcuts disarankan menggunakan kata sandi admin yang kuat, HTTPS, pencadangan basis data dan media secara berkala, pemantauan log, serta pemisahan akun sesuai tanggung jawab apabila jumlah petugas bertambah.",
        "Pengembangan selanjutnya dapat menambahkan pencetakan struk, ekspor laporan, audit perubahan transaksi, manajemen pengembalian dana, dan peran pengguna yang lebih rinci. Integrasi pembayaran digital dapat dipertimbangkan setelah alur rekonsiliasi, keamanan callback, dan penanganan transaksi gagal dirancang serta diuji secara memadai.",
        "Pengujian lanjutan disarankan mencakup kompatibilitas berbagai perangkat, beban pengguna bersamaan, pemulihan setelah gangguan jaringan, serta uji penerimaan pada operasional Homcuts yang sebenarnya.",
    ]
    for i, item in enumerate(suggestions, 1):
        base.add_list_item(doc, f"{i})", item)


def add_bibliography(doc: Document) -> None:
    doc.add_page_break()
    add_centered_title(doc, "DAFTAR PUSTAKA")
    refs = [
        "Apriandi, M. N., Irawan, A. S. Y., & Purwantoro. (2025). Implementasi framework Laravel pada aplikasi pemesanan barbershop berbasis web (studi kasus: Maiden Barberrock). Jurnal Informatika dan Teknik Elektro Terapan, 13(3S1), 96–108. https://doi.org/10.23960/jitet.v13i3S1.7520",
        "Bangor, A., Kortum, P., & Miller, J. (2009). Determining what individual SUS scores mean: Adding an adjective rating scale. Journal of Usability Studies, 4(3), 114–123. https://uxpajournal.org/wp-content/uploads/sites/7/pdf/JUS_Bangor_May2009.pdf",
        "Brooke, J. (1996). SUS: A ‘quick and dirty’ usability scale. Dalam P. W. Jordan, B. Thomas, A. Weerdmeester, & I. L. McClelland (Ed.), Usability Evaluation in Industry. Taylor & Francis. https://hci-studies.org/methods-and-measures/downloads/SUS_Brooke1996.pdf",
        "Docker. (2026). Docker Compose documentation. https://docs.docker.com/compose/",
        "Firmansyah, D., Purwanto, H., Wiharko, T., & Purbayanto, B. (2023). Aplikasi booking barbershop online berbasis web. Jurnal Internal, 6(2), 146–155. https://jurnal.masoemuniversity.ac.id/index.php/internal/article/download/849/596",
        "Ginoga, W. S., Hidayat, M., & Pakaya, N. (2023). Sistem informasi akuntansi barbershop. Journal of System and Information Technology, 3(1), 27–36. https://ejurnal.ung.ac.id/index.php/diffusion/article/download/12408/5963",
        "International Organization for Standardization. (2018). ISO 9241-11:2018 Ergonomics of human-system interaction—Part 11: Usability: Definitions and concepts. https://www.iso.org/standard/63500.html",
        "International Software Testing Qualifications Board. (2024). Certified Tester Foundation Level syllabus v4.0.1. https://istqb.org/wp-content/uploads/2024/11/ISTQB_CTFL_Syllabus_v4.0.1.pdf",
        "Laravel. (2026). Laravel 13.x documentation. https://laravel.com/docs/13.x",
        "Laudon, K. C., & Laudon, J. P. (2022). Management Information Systems: Managing the Digital Firm (17th ed.). Pearson.",
        "Ngatini, N., & Cahyanti, F. L. D. (2024). Perancangan sistem point of sales (POS) berbasis web untuk optimalisasi layanan pada Shortcut Barbershop. Jurnal Nasional Komputasi dan Teknologi Informasi, 7(6), 1707–1715. https://doi.org/10.32672/jnkti.v7i6.8260",
        "Object Management Group. (2017). OMG Unified Modeling Language (OMG UML), version 2.5.1. https://www.omg.org/spec/UML/2.5.1/PDF",
        "Oracle. (2026). MySQL 8.4 Reference Manual. https://dev.mysql.com/doc/refman/8.4/en/",
        "Sommerville, I. (2016). Software Engineering (10th ed.). Pearson.",
    ]
    for ref in refs:
        base.add_bibliography_entry(doc, ref)


def set_footer_and_numbering(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(4)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(4)
        section.right_margin = Cm(3)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.5)
        base.set_page_number_start(section, 1)
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        field = p.add_run()
        base.style_run(field, size=12)
        base.add_field_code(field)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Dokumen revisi sumber tidak ditemukan: {SOURCE}")
    doc = Document(SOURCE)
    remove_old_bibliography(doc)
    prepend_bab_i(doc)
    add_bab_iv(doc)
    add_bab_v(doc)
    add_bibliography(doc)
    normalize_chapter_breaks(doc)
    set_footer_and_numbering(doc)
    doc.core_properties.title = "Sistem Informasi Manajemen Homcuts — BAB I sampai BAB V"
    doc.core_properties.subject = "Naskah tugas akhir sesuai pedoman D4 PNUP"
    doc.core_properties.author = "Draf akademik penelitian Homcuts"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
