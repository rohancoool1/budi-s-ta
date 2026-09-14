from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image, ImageEnhance, ImageFilter
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/home/rogan/Documents/budi/da-project")
OUT_DIR = ROOT / "output" / "docx"
OUT_PATH = OUT_DIR / "Revisi_BAB_II_dan_BAB_III_Homcuts.docx"
WORK = Path("/tmp/revise_bab23/generated")
SOURCE_240 = Path("/tmp/revise_bab23/source_pages_240")


def run_dot(name: str, source: str) -> Path:
    dot_path = WORK / f"{name}.dot"
    png_path = WORK / f"{name}.png"
    dot_path.write_text(source, encoding="utf-8")
    subprocess.run(
        ["dot", "-Tpng", "-Gdpi=220", str(dot_path), "-o", str(png_path)],
        check=True,
    )
    return png_path


def crop_interface(source: Path, box: tuple[int, int, int, int], name: str) -> Path:
    image = Image.open(source).convert("RGB")
    cropped = image.crop(box)
    cropped = ImageEnhance.Contrast(cropped).enhance(1.05)
    cropped = ImageEnhance.Sharpness(cropped).enhance(1.2)
    cropped = cropped.filter(ImageFilter.UnsharpMask(radius=1.2, percent=115, threshold=3))
    target = WORK / name
    cropped.save(target, quality=95, dpi=(240, 240))
    return target


def make_diagrams() -> dict[str, Path]:
    diagram: dict[str, Path] = {}
    common = 'graph [bgcolor="white", pad="0.18", nodesep="0.35", ranksep="0.5"]; node [fontname="Arial", fontsize=11, color="#303030", penwidth=1.2]; edge [fontname="Arial", fontsize=9, color="#555555", penwidth=1.1, arrowsize=0.75];'

    diagram["procedure"] = run_dot(
        "procedure",
        f'''digraph G {{ rankdir=LR; {common}
        node [shape=box, style="rounded,filled", fillcolor="#F8F5EF", margin="0.16,0.10"];
        a [label="Identifikasi kebutuhan"];
        b [label="Pemodelan dan\nperancangan"];
        c [label="Implementasi\nsistem"];
        d [label="Pengujian\nsistem"];
        e [label="Evaluasi dan\nperbaikan", fillcolor="#E8EFE2"];
        a -> b -> c -> d -> e;
        e -> b [label="umpan balik", style=dashed, constraint=false, color="#748164"];
        }}''',
    )

    diagram["use_case"] = run_dot(
        "use_case",
        f'''digraph G {{ rankdir=LR; {common}
        node [shape=ellipse, style="filled", fillcolor="#FBFAF7", margin="0.18,0.10"];
        pelanggan [shape=box, label="Pelanggan", style="rounded,filled", fillcolor="#E8EFE2"];
        admin [shape=box, label="Admin / Kasir", style="rounded,filled", fillcolor="#FCE7DC"];
        subgraph cluster_system {{ label="Sistem Informasi Manajemen Homcuts"; color="#777777"; style="rounded";
          subgraph public_cases {{ rank=same;
            info [label="Mengakses informasi\nHomcuts"];
            booking [label="Membuat booking\nlayanan"];
            product [label="Memesan produk"];
            status [label="Melihat status booking\natau pesanan"];
            contact [label="Mengirim pesan"];
          }}
          subgraph admin_cases {{ rank=same;
            login [label="Login admin"];
            pos [label="Mencatat transaksi\nwalk-in melalui POS"];
            confirm [label="Mengonfirmasi\npembayaran tunai"];
            manage [label="Mengelola data dan\naktivitas operasional"];
            summary [label="Melihat ringkasan\noperasional dan transaksi"];
          }}
          info -> login [style=invis, weight=20];
        }}
        pelanggan -> info [dir=none]; pelanggan -> booking [dir=none]; pelanggan -> product [dir=none]; pelanggan -> status [dir=none]; pelanggan -> contact [dir=none];
        login -> admin [dir=none]; pos -> admin [dir=none]; confirm -> admin [dir=none]; manage -> admin [dir=none]; summary -> admin [dir=none];
        }}''',
    )

    diagram["activity_booking"] = run_dot(
        "activity_booking",
        f'''digraph G {{ rankdir=TB; {common}
        node [shape=box, style="rounded,filled", fillcolor="#FBFAF7", margin="0.16,0.08"];
        start [shape=circle, label="", width=0.18, fixedsize=true, fillcolor="#222222"];
        form [label="Pelanggan memilih layanan, barber,\ntanggal, dan waktu"];
        validate [label="Sistem memvalidasi data, jam kerja,\ndurasi, dan benturan jadwal"];
        ok [shape=diamond, label="Slot tersedia?", fillcolor="#FFF4CE"];
        reject [label="Tolak dan tampilkan\npesan slot tidak tersedia", fillcolor="#FCE7DC"];
        save [label="Simpan booking dan transaksi;\nbuat nomor antrean terpadu"];
        wait [label="Status: menunggu\npembayaran tunai"];
        edit [label="Admin dapat mengubah jadwal;\ninformasi pelanggan ikut diperbarui", fillcolor="#E8EFE2"];
        paid [label="Kasir mengonfirmasi uang diterima\ndi halaman Booking atau Transaksi"];
        service [label="Layanan dilaksanakan\ndan diselesaikan"];
        end [shape=doublecircle, label="", width=0.18, fixedsize=true, fillcolor="#222222"];
        start -> form -> validate -> ok;
        ok -> reject [label="tidak"]; reject -> form;
        ok -> save [label="ya"]; save -> wait -> paid -> service -> end;
        wait -> edit [label="bila dijadwal ulang"]; edit -> wait;
        }}''',
    )

    diagram["activity_pos"] = run_dot(
        "activity_pos",
        f'''digraph G {{ rankdir=TB; {common}
        node [shape=box, style="rounded,filled", fillcolor="#FBFAF7", margin="0.16,0.08"];
        start [shape=circle, label="", width=0.18, fixedsize=true, fillcolor="#222222"];
        input [label="Kasir memasukkan pelanggan, layanan,\nbarber, dan waktu layanan"];
        validate [label="Sistem memeriksa jam kerja dan benturan\ndengan booking maupun walk-in lain"];
        ok [shape=diamond, label="Jadwal tersedia?", fillcolor="#FFF4CE"];
        reject [label="Tolak dan tampilkan\njadwal yang bentrok", fillcolor="#FCE7DC"];
        save [label="Simpan transaksi POS dan berikan\nnomor antrean terpadu"];
        queue [label="Pelanggan menunggu lalu\nmenerima layanan"];
        paid [label="Setelah layanan, kasir mengonfirmasi\npembayaran tunai"];
        finish [label="Status transaksi menjadi selesai"];
        end [shape=doublecircle, label="", width=0.18, fixedsize=true, fillcolor="#222222"];
        start -> input -> validate -> ok;
        ok -> reject [label="tidak"]; reject -> input;
        ok -> save [label="ya"]; save -> queue -> paid -> finish -> end;
        }}''',
    )

    diagram["activity_product"] = run_dot(
        "activity_product",
        f'''digraph G {{ rankdir=TB; {common}
        node [shape=box, style="rounded,filled", fillcolor="#FBFAF7", margin="0.16,0.08"];
        start [shape=circle, label="", width=0.18, fixedsize=true, fillcolor="#222222"];
        cart [label="Pelanggan memilih produk dan\nmenambahkannya ke keranjang"];
        checkout [label="Pelanggan mengisi data checkout"];
        check [label="Sistem memvalidasi data, menghitung ulang\ntotal, dan memeriksa stok"];
        ok [shape=diamond, label="Stok cukup?", fillcolor="#FFF4CE"];
        reject [label="Tolak pesanan dan tampilkan\ninformasi stok", fillcolor="#FCE7DC"];
        save [label="Simpan transaksi, item, pembayaran tunai,\ndan kurangi stok"];
        wait [label="Status: menunggu pembayaran"];
        paid [label="Kasir mengonfirmasi\npembayaran di halaman Transaksi"];
        ready [label="Pesanan siap diambil"];
        end [shape=doublecircle, label="", width=0.18, fixedsize=true, fillcolor="#222222"];
        start -> cart -> checkout -> check -> ok;
        ok -> reject [label="tidak"]; reject -> cart;
        ok -> save [label="ya"]; save -> wait -> paid -> ready -> end;
        }}''',
    )

    diagram["erd"] = run_dot(
        "erd",
        f'''digraph G {{ rankdir=TB; {common} graph [bgcolor="white", pad="0.1", nodesep="0.24", ranksep="0.42"];
        node [shape=record, style="filled", fillcolor="#FBFAF7", fontsize=9.5, margin="0.07"];
        users [label="{{USERS|PK id\\l|name\\lemail\\lpassword\\lis_admin\\l}}"];
        barbers [label="{{BARBERS|PK id\\l|name\\lwork_start_time\\lwork_end_time\\limage_path\\lis_active\\l}}"];
        services [label="{{SERVICES|PK id\\l|name\\lduration_minutes\\lprice\\lis_active\\l}}"];
        products [label="{{PRODUCTS|PK id\\l|name\\lprice\\lstock\\limage_path\\lis_active\\l}}"];
        bookings [label="{{BOOKINGS|PK id\\lFK barber_id\\lFK service_catalog_id\\l|starts_at\\lends_at\\lduration_minutes\\lstatus\\l}}"];
        orders [label="{{ORDERS|PK id\\lFK booking_id\\lFK cashier_id\\l|channel\\ltransaction_type\\lqueue_number\\lservice_starts_at\\lservice_ends_at\\lpayment_status\\ltotal\\l}}"];
        items [label="{{ORDER_ITEMS|PK id\\lFK order_id\\lFK product_id\\lFK service_id\\lFK barber_id\\l|item_type\\lunit_price\\lquantity\\lline_total\\l}}"];
        payments [label="{{PAYMENTS|PK id\\lFK order_id\\lFK confirmed_by\\l|method\\lstatus\\lamount\\lpaid_at\\l}}"];
        gallery [label="{{GALLERY_ENTRIES|PK id\\lFK barber_id\\l|style\\lquote\\limage_path\\lis_published\\l}}"];
        sequences [label="{{QUEUE_SEQUENCES|PK name\\l|last_number\\l}}"];
        {{rank=same; barbers; services; products; users;}}
        {{rank=same; bookings; gallery; orders;}}
        {{rank=same; items; payments; sequences;}}
        barbers -> bookings [label="1 : N"];
        services -> bookings [label="1 : N"];
        bookings -> orders [label="1 : 0..1"];
        users -> orders [label="1 : N\n(kasir)"];
        orders -> items [label="1 : N"];
        products -> items [label="1 : N"];
        services -> items [label="1 : N"];
        barbers -> items [label="1 : N"];
        orders -> payments [label="1 : N"];
        users -> payments [label="1 : N\n(konfirmasi)"];
        barbers -> gallery [label="1 : N"];
        sequences -> orders [label="menghasilkan\nnomor antrean", style=dashed];
        }}''',
    )

    diagram["navigation"] = run_dot(
        "navigation",
        f'''digraph G {{ rankdir=TB; {common}
        node [shape=box, style="rounded,filled", fillcolor="#FBFAF7", margin="0.14,0.08"];
        root [label="Sistem Homcuts", fillcolor="#E8EFE2"];
        public [label="Halaman Pelanggan", fillcolor="#F3EFE6"];
        admin [label="Halaman Admin / Kasir", fillcolor="#FCE7DC"];
        root -> public; root -> admin;
        public -> home [label=""]; home [label="Beranda"];
        public -> booking [label=""]; booking [label="Booking"];
        public -> shop [label=""]; shop [label="Toko Produk"];
        public -> gallery [label=""]; gallery [label="Galeri"];
        public -> about [label=""]; about [label="Tentang Kami"];
        public -> contact [label=""]; contact [label="Kontak"];
        public -> status [label=""]; status [label="Status Transaksi"];
        admin -> dashboard [label=""]; dashboard [label="Ringkasan"];
        admin -> pos [label=""]; pos [label="Kasir POS"];
        admin -> admbook [label=""]; admbook [label="Booking"];
        admin -> trans [label=""]; trans [label="Transaksi"];
        admin -> master [label=""]; master [label="Barber · Layanan · Produk · Galeri"];
        admin -> other [label=""]; other [label="Pesan · Pengaturan · Akun"];
        }}''',
    )

    diagram["architecture"] = run_dot(
        "architecture",
        f'''digraph G {{ rankdir=TB; {common}
        node [shape=box, style="rounded,filled", margin="0.18,0.10"];
        browser [label="Peramban Pelanggan / Admin\nHTML · CSS · JavaScript", fillcolor="#E8EFE2"];
        vite [label="Aset antarmuka\nBlade · Tailwind CSS · Vite", fillcolor="#F3EFE6"];
        routes [label="Laravel Web Layer\nRoute · Middleware · Controller", fillcolor="#FBFAF7"];
        service [label="Lapisan Logika Bisnis\nAvailability · Queue · Transaction · Payment · Notification", fillcolor="#FFF4CE"];
        model [label="Model Eloquent / Query Builder", fillcolor="#FBFAF7"];
        mysql [label="MySQL 8.4\nData transaksi dan data utama", fillcolor="#FCE7DC"];
        storage [label="Laravel Storage\nFoto produk, barber, dan galeri", fillcolor="#FCE7DC"];
        scheduler [label="Laravel Scheduler\nPemrosesan berkala", fillcolor="#E8EFE2"];
        docker [shape=folder, label="Docker Compose: web · db · scheduler\nvolume basis data dan media", fillcolor="#EAF2F8"];
        browser -> vite -> routes -> service -> model -> mysql;
        service -> storage;
        scheduler -> service;
        docker -> routes [style=dashed]; docker -> mysql [style=dashed]; docker -> scheduler [style=dashed];
        }}''',
    )

    diagram["ui_customer"] = crop_interface(
        SOURCE_240 / "page-27.png",
        (375, 760, 1705, 1520),
        "ui_customer.jpg",
    )
    diagram["ui_admin"] = crop_interface(
        SOURCE_240 / "page-28.png",
        (375, 575, 1705, 1355),
        "ui_admin.jpg",
    )
    return diagram


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size: float = 9.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(caption)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, header in enumerate(headers):
        set_cell_text(header_row.cells[idx], header, bold=True, size=font_size)
        set_cell_shading(header_row.cells[idx], "E7E3DB")
        if widths:
            set_cell_width(header_row.cells[idx], widths[idx])

    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            set_cell_text(row.cells[idx], value, size=font_size)
            if widths:
                set_cell_width(row.cells[idx], widths[idx])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    return table


def add_field_code(field) -> None:
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    field._r.append(fld_char1)
    field._r.append(instr_text)
    field._r.append(fld_char2)


def set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def style_run(run, bold: bool = False, italic: bool = False, size: float = 12) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_rich_paragraph(doc: Document, text: str, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, italic_terms: list[str] | None = None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(1.27) if indent else Cm(0)
    italic_terms = sorted(italic_terms or [], key=len, reverse=True)
    if not italic_terms:
        style_run(p.add_run(text))
        return p

    pattern = "(" + "|".join(re.escape(term) for term in italic_terms) + ")"
    for part in re.split(pattern, text):
        if not part:
            continue
        style_run(p.add_run(part), italic=part in italic_terms)
    return p


def add_list_item(doc: Document, marker: str, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-0.7)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run(f"{marker} "), bold=False)
    style_run(p.add_run(text))
    return p


def add_section_heading(doc: Document, number: str, title: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{number} {title}")
    style_run(run, bold=(level == 1))
    return p


def add_chapter_heading(doc: Document, chapter: str, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(chapter)
    style_run(run, bold=True)
    run.add_break()
    run2 = p.add_run(title)
    style_run(run2, bold=True)
    return p


def add_figure(doc: Document, path: Path, caption: str, width_cm: float = 14.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Cm(0)
    cp.paragraph_format.line_spacing = 1.0
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(6)
    cp.paragraph_format.keep_with_next = True
    style_run(cp.add_run(caption), size=11)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.5)
    set_page_number_start(section, 4)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.first_line_indent = Cm(1.27)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    field = p.add_run()
    style_run(field, size=12)
    add_field_code(field)

    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:space"), "720")


def add_bibliography_entry(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.space_after = Pt(6)
    style_run(p.add_run(text), size=11)


def build_document(diagram: dict[str, Path]) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Revisi BAB II dan BAB III Sistem Informasi Manajemen Homcuts"
    doc.core_properties.subject = "Tinjauan Pustaka dan Metode Penelitian"
    doc.core_properties.author = "Draf akademik untuk penelitian Homcuts"

    # BAB II
    add_chapter_heading(doc, "BAB II", "TINJAUAN PUSTAKA")
    add_rich_paragraph(
        doc,
        "Bab ini menguraikan konsep yang menjadi dasar perancangan, implementasi, dan pengujian sistem informasi manajemen Homcuts. Pembahasan disusun dari konsep umum sistem informasi hingga konsep khusus yang berkaitan dengan booking layanan, antrean, transaksi kasir, penjualan produk, teknologi pengembangan, pemodelan, dan pengujian sistem.",
        italic_terms=["booking"],
    )

    add_section_heading(doc, "2.1", "Sistem Informasi Manajemen")
    add_rich_paragraph(
        doc,
        "Sistem informasi merupakan sekumpulan komponen yang saling berhubungan untuk mengumpulkan, mengolah, menyimpan, dan menyajikan informasi. Komponen tersebut mencakup manusia, prosedur, perangkat lunak, perangkat keras, basis data, dan jaringan. Apabila diarahkan untuk mendukung kegiatan operasional serta pengambilan keputusan, sistem tersebut disebut sistem informasi manajemen (Laudon & Laudon, 2022).",
    )
    add_rich_paragraph(
        doc,
        "Nilai utama sistem informasi manajemen tidak hanya terletak pada digitalisasi pencatatan, tetapi juga pada integrasi data antarkegiatan. Data yang dicatat pada satu proses dapat digunakan kembali pada proses lain sehingga duplikasi, ketidakkonsistenan, dan keterlambatan informasi dapat dikurangi. Dalam penelitian ini, integrasi diwujudkan melalui hubungan antara data pelanggan, jadwal layanan, barber, produk, antrean, transaksi, dan pembayaran.",
        italic_terms=["barber"],
    )

    add_section_heading(doc, "2.2", "Sistem Informasi Manajemen Barbershop")
    add_rich_paragraph(
        doc,
        "Sistem informasi manajemen barbershop adalah aplikasi yang mengelola informasi publik dan aktivitas operasional usaha jasa pangkas rambut. Informasi publik meliputi profil usaha, daftar layanan dan harga, barber, galeri hasil potongan, produk, serta kontak. Aktivitas operasional meliputi booking, penjadwalan, pelayanan pelanggan walk-in, penjualan produk, pembayaran, dan pelaporan transaksi.",
        italic_terms=["barbershop", "barber", "booking", "walk-in"],
    )
    add_rich_paragraph(
        doc,
        "Karakteristik utama usaha barbershop adalah keterikatan layanan pada sumber daya tertentu, yaitu barber dan waktu. Satu barber tidak dapat melayani dua pelanggan pada selang waktu yang sama. Oleh karena itu, sistem harus mampu menggabungkan informasi durasi layanan, jadwal kerja barber, sumber transaksi, serta status pelayanan untuk menghasilkan jadwal yang konsisten.",
        italic_terms=["barbershop", "barber"],
    )

    add_section_heading(doc, "2.3", "Pemesanan dan Penjadwalan Layanan")
    add_rich_paragraph(
        doc,
        "Booking online memungkinkan pelanggan memilih layanan dan waktu kunjungan sebelum datang ke lokasi. Informasi yang dibutuhkan sekurang-kurangnya meliputi identitas pelanggan, layanan, barber pilihan, tanggal, waktu mulai, dan durasi. Sistem kemudian memeriksa apakah permintaan berada dalam jam operasional dan jam kerja barber serta tidak bertabrakan dengan jadwal yang telah tercatat.",
        italic_terms=["Booking online", "barber"],
    )
    add_rich_paragraph(
        doc,
        "Benturan jadwal dapat ditentukan dengan membandingkan interval waktu. Dua layanan A dan B bertabrakan apabila waktu mulai A lebih kecil daripada waktu selesai B dan waktu selesai A lebih besar daripada waktu mulai B. Pemeriksaan berbasis interval lebih tepat daripada membandingkan jam mulai saja karena setiap layanan dapat memiliki durasi yang berbeda. Pada Homcuts, waktu layanan berada dalam rentang operasional pukul 07.00 sampai 22.00; waktu mulai terakhir mengikuti durasi layanan dan batas akhir kerja barber.",
        italic_terms=["barber"],
    )
    add_rich_paragraph(
        doc,
        "Perubahan jadwal oleh administrator harus memperbarui sumber data yang sama dengan yang dibaca pelanggan. Dengan demikian, informasi tanggal, waktu, barber, dan status yang ditampilkan kepada pelanggan selalu mengacu pada data terbaru, bukan salinan yang terpisah.",
        italic_terms=["barber"],
    )

    add_section_heading(doc, "2.4", "Sistem Antrean Pelayanan")
    add_rich_paragraph(
        doc,
        "Sistem antrean mengatur identitas dan urutan pelanggan yang akan menerima layanan. Nomor antrean harus bersifat unik dalam ruang lingkup yang ditetapkan agar tidak terjadi dua pelanggan dengan nomor yang sama. Pada sistem Homcuts, booking yang dibuat pelanggan dan transaksi walk-in yang dimasukkan kasir menggunakan sumber penomoran yang sama. Dengan cara ini, perbedaan saluran pencatatan tidak menghasilkan dua antrean yang saling terpisah.",
        italic_terms=["booking", "walk-in"],
    )
    add_rich_paragraph(
        doc,
        "Nomor antrean berfungsi sebagai pengenal pelayanan, sedangkan waktu mulai dan waktu selesai berfungsi sebagai dasar ketersediaan jadwal. Pemisahan fungsi tersebut penting: nomor yang lebih kecil tidak selalu berarti pelanggan bebas memilih waktu yang telah digunakan, dan perubahan jadwal tidak boleh menghasilkan nomor duplikat.",
    )

    add_section_heading(doc, "2.5", "Point of Sale dan Pengelolaan Transaksi Tunai")
    add_rich_paragraph(
        doc,
        "Point of Sale (POS) adalah bagian sistem yang digunakan kasir untuk mencatat transaksi pada titik pelayanan. Dalam barbershop, POS tidak hanya menangani produk, tetapi juga layanan pelanggan walk-in. Kasir memasukkan pelanggan, layanan, barber, waktu, produk tambahan bila ada, dan nilai transaksi. Ngatini dan Cahyanti (2024) menunjukkan bahwa POS berbasis web dapat memusatkan pencatatan transaksi dan membantu pengelolaan operasional barbershop.",
        italic_terms=["Point of Sale", "walk-in", "barber", "barbershop"],
    )
    add_rich_paragraph(
        doc,
        "Status transaksi dan status pembayaran merupakan dua informasi yang berbeda. Status transaksi menunjukkan tahap pemenuhan layanan atau pesanan, misalnya menunggu, siap diambil, selesai, atau dibatalkan. Status pembayaran menunjukkan apakah uang telah diterima, misalnya belum dibayar atau lunas. Pemisahan ini diperlukan karena layanan dapat selesai tetapi belum dibayar, sedangkan pesanan produk dapat sudah dibayar namun masih menunggu untuk diambil.",
    )
    add_rich_paragraph(
        doc,
        "Penelitian ini menggunakan pembayaran tunai secara manual. Data transaksi dibuat dengan status belum dibayar, kemudian kasir mengonfirmasi setelah uang diterima. Konfirmasi booking dapat dilakukan melalui halaman Booking atau Transaksi, sedangkan pembayaran pesanan produk dikonfirmasi melalui halaman Transaksi. Transaksi yang telah dikonfirmasi menjadi sumber laporan penerimaan usaha.",
        italic_terms=["booking"],
    )

    add_section_heading(doc, "2.6", "Penjualan Produk dan Pengelolaan Persediaan")
    add_rich_paragraph(
        doc,
        "Fitur penjualan produk memungkinkan pelanggan memilih barang, menambahkannya ke keranjang, dan melakukan checkout untuk pengambilan langsung di barbershop. Sistem perlu menghitung total berdasarkan harga yang tersimpan pada server, bukan hanya nilai yang dikirim peramban, agar perubahan nilai pada sisi klien tidak mengubah jumlah tagihan.",
        italic_terms=["checkout", "barbershop"],
    )
    add_rich_paragraph(
        doc,
        "Persediaan harus diperiksa pada saat transaksi dibuat. Jika stok mencukupi, jumlah barang yang dipesan dicatat pada item transaksi dan stok dikurangi atau ditahan sesuai aturan sistem. Jika transaksi dibatalkan sebelum diselesaikan, stok perlu dikembalikan secara konsisten. Relasi antara transaksi, item transaksi, dan produk memungkinkan setiap perubahan stok ditelusuri ke sumbernya.",
    )

    add_section_heading(doc, "2.7", "Aplikasi Web dan Arsitektur Model-View-Controller")
    add_rich_paragraph(
        doc,
        "Aplikasi web dijalankan melalui peramban dan berkomunikasi dengan server menggunakan permintaan HTTP. Pendekatan ini memungkinkan pelanggan dan administrator mengakses sistem melalui perangkat yang berbeda tanpa memasang aplikasi khusus. Antarmuka pelanggan dan antarmuka admin tetap menggunakan sumber data yang sama pada server.",
    )
    add_rich_paragraph(
        doc,
        "Arsitektur Model-View-Controller (MVC) memisahkan pengelolaan data, tampilan, dan pengendalian alur permintaan. Model merepresentasikan data dan relasi basis data, view menyajikan antarmuka, sedangkan controller menerima permintaan, melakukan validasi, memanggil logika bisnis, dan menentukan respons. Pemisahan ini meningkatkan keteraturan kode dan memudahkan pengujian serta pemeliharaan (Sommerville, 2016).",
        italic_terms=["Model-View-Controller", "Model", "view", "controller"],
    )

    add_section_heading(doc, "2.8", "Teknologi Pengembangan Sistem")
    add_section_heading(doc, "2.8.1", "PHP dan Laravel", level=2)
    add_rich_paragraph(
        doc,
        "PHP merupakan bahasa pemrograman sisi server yang digunakan untuk memproses permintaan, menjalankan aturan bisnis, dan menghasilkan respons web. Laravel menyediakan struktur MVC, routing, middleware, validasi, autentikasi, Eloquent ORM, migrasi basis data, penyimpanan berkas, penjadwal tugas, serta fasilitas pengujian. Fitur tersebut digunakan agar pengembangan sistem Homcuts terstruktur dan konsisten (Laravel, 2026).",
        italic_terms=["routing", "middleware"],
    )

    add_section_heading(doc, "2.8.2", "Blade, Tailwind CSS, JavaScript, dan Vite", level=2)
    add_rich_paragraph(
        doc,
        "Blade digunakan sebagai mesin templat untuk membentuk halaman dari data Laravel. Tailwind CSS menyediakan kelas utilitas untuk menyusun tampilan responsif, sedangkan JavaScript menangani interaksi pada peramban, seperti keranjang, pemeriksaan ketersediaan jadwal, pengurutan tabel, notifikasi, dan pembaruan status tanpa muat ulang manual. Vite digunakan pada tahap pengembangan untuk membangun dan mengoptimalkan aset CSS serta JavaScript. Vite bukan kerangka antarmuka, sehingga penggunaan teknologi ini tidak mewajibkan React atau TypeScript.",
        italic_terms=["Blade", "JavaScript", "Vite", "React", "TypeScript"],
    )

    add_section_heading(doc, "2.8.3", "MySQL", level=2)
    add_rich_paragraph(
        doc,
        "MySQL adalah sistem manajemen basis data relasional. Data disimpan dalam tabel yang dihubungkan melalui kunci primer dan kunci asing. Transaksi basis data digunakan untuk memastikan beberapa perubahan yang saling berkaitan, misalnya pembuatan pesanan, item, pembayaran, nomor antrean, dan perubahan stok, diselesaikan sebagai satu kesatuan. Penguncian data saat proses kritis membantu mencegah dua permintaan menggunakan slot atau stok yang sama (Oracle, 2026).",
    )

    add_section_heading(doc, "2.8.4", "Docker dan Docker Compose", level=2)
    add_rich_paragraph(
        doc,
        "Docker mengemas aplikasi dan dependensinya ke dalam container. Docker Compose mendefinisikan beberapa layanan yang bekerja bersama melalui satu konfigurasi, misalnya aplikasi web, basis data MySQL, dan scheduler. Volume digunakan untuk mempertahankan data basis data serta berkas unggahan ketika container dibuat ulang. Pendekatan ini menghasilkan lingkungan pengembangan yang lebih seragam (Docker, 2026).",
        italic_terms=["container", "scheduler"],
    )

    add_section_heading(doc, "2.9", "Pemodelan dan Perancangan Sistem")
    add_section_heading(doc, "2.9.1", "Use Case Diagram", level=2)
    add_rich_paragraph(
        doc,
        "Use case diagram menggambarkan fungsi yang disediakan sistem dan hubungan fungsi tersebut dengan aktor. Diagram ini digunakan untuk menetapkan batas sistem dan tanggung jawab Pelanggan serta Admin/Kasir. Hubungan aktor dengan use case menunjukkan interaksi, bukan urutan proses (Object Management Group, 2017).",
        italic_terms=["Use case diagram", "use case"],
    )

    add_section_heading(doc, "2.9.2", "Activity Diagram", level=2)
    add_rich_paragraph(
        doc,
        "Activity diagram menggambarkan aliran aktivitas, keputusan, dan kemungkinan pengulangan pada suatu proses. Dalam penelitian ini, activity diagram digunakan untuk menjelaskan alur booking, transaksi walk-in melalui POS, dan pemesanan produk. Diagram memisahkan tindakan pengguna, proses sistem, dan tindakan admin/kasir agar tanggung jawab setiap pihak terlihat jelas.",
        italic_terms=["Activity diagram", "booking", "walk-in"],
    )

    add_section_heading(doc, "2.9.3", "Entity Relationship Diagram", level=2)
    add_rich_paragraph(
        doc,
        "Entity Relationship Diagram (ERD) menggambarkan entitas, atribut penting, kunci, dan hubungan antarentitas dalam basis data. ERD digunakan untuk memastikan bahwa data utama dan data transaksi tidak disimpan secara berulang serta bahwa hubungan antara booking, pesanan, item, pembayaran, layanan, barber, dan produk dapat ditelusuri.",
        italic_terms=["Entity Relationship Diagram", "booking", "barber"],
    )

    add_section_heading(doc, "2.9.4", "Struktur Navigasi, Arsitektur, dan Antarmuka", level=2)
    add_rich_paragraph(
        doc,
        "Struktur navigasi menunjukkan hubungan antarhalaman yang dapat diakses pengguna. Rancangan arsitektur menunjukkan pembagian antarmuka, lapisan aplikasi, logika bisnis, penyimpanan, dan basis data. Rancangan antarmuka menerjemahkan kebutuhan menjadi susunan layar sebelum implementasi. Ketiganya digunakan untuk menjaga konsistensi antara kebutuhan fungsional, proses, dan tampilan.",
    )

    add_section_heading(doc, "2.10", "Pengujian Sistem")
    add_rich_paragraph(
        doc,
        "Pengujian dilakukan untuk menemukan ketidaksesuaian antara perilaku sistem dan kebutuhan yang telah ditetapkan. Landasan pengujian ditempatkan pada Bab II, rancangan dan langkah pengujiannya dijelaskan pada Bab III, sedangkan hasil pelaksanaan serta pembahasannya disajikan pada Bab IV.",
    )

    add_section_heading(doc, "2.10.1", "Black-Box Testing", level=2)
    add_rich_paragraph(
        doc,
        "Black-box testing merupakan pengujian berbasis spesifikasi yang menilai masukan dan keluaran tanpa bergantung pada struktur internal program. Kasus uji diturunkan dari kebutuhan, aturan validasi, dan alur bisnis. Teknik ini sesuai untuk memeriksa fungsi halaman publik, booking, pemeriksaan jadwal, POS, checkout produk, autentikasi admin, konfirmasi pembayaran, pengelolaan data, dan pembaruan status (ISTQB, 2024).",
        italic_terms=["Black-box testing", "booking", "checkout"],
    )
    add_rich_paragraph(
        doc,
        "Setiap kasus uji memuat skenario, data masukan, hasil yang diharapkan, hasil aktual, dan status lulus atau gagal. Persentase keberhasilan fungsional dihitung menggunakan rumus: jumlah kasus uji lulus dibagi jumlah seluruh kasus uji, kemudian dikalikan 100 persen.",
    )

    add_section_heading(doc, "2.10.2", "User Acceptance Testing", level=2)
    add_rich_paragraph(
        doc,
        "User Acceptance Testing (UAT) merupakan pengujian penerimaan yang dilakukan oleh pengguna atau perwakilan pemilik proses bisnis untuk memastikan sistem dapat digunakan sesuai kebutuhan operasional. UAT menilai sistem dari sudut pandang pekerjaan nyata, bukan struktur internal program. Skenario UAT diturunkan dari kebutuhan pengguna dan memuat tugas, hasil yang diharapkan, keputusan diterima atau tidak diterima, serta catatan pengguna. Pengujian penerimaan dilakukan setelah fungsi utama dinyatakan stabil melalui pengujian fungsional (ISTQB, 2024; Sommerville, 2016).",
        italic_terms=["User Acceptance Testing", "UAT"],
    )
    add_rich_paragraph(
        doc,
        "Black-box testing dan UAT sama-sama menggunakan masukan serta keluaran yang dapat diamati, tetapi keduanya memiliki tujuan dan pelaksana yang berbeda. Black-box testing memeriksa kesesuaian fungsi terhadap spesifikasi, sedangkan UAT menentukan apakah alur tersebut dapat diterima untuk pekerjaan pengguna. Pada sistem Homcuts, UAT perlu mewakili dua aktor, yaitu Pelanggan serta Admin/Kasir. Persentase penerimaan dihitung dengan membagi jumlah skenario yang diterima dengan seluruh skenario yang dijalankan, kemudian dikalikan 100 persen. Skenario kritis, seperti pencegahan benturan jadwal dan konfirmasi pembayaran, harus diterima seluruhnya sebelum sistem dinyatakan layak digunakan.",
        italic_terms=["Black-box testing", "UAT"],
    )

    add_section_heading(doc, "2.10.3", "System Usability Scale", level=2)
    add_rich_paragraph(
        doc,
        "System Usability Scale (SUS) adalah instrumen berisi sepuluh pernyataan dengan skala jawaban 1 sampai 5 untuk memperoleh ukuran kegunaan sistem secara ringkas (Brooke, 1996). Pernyataan bernomor ganjil bersifat positif dan pernyataan bernomor genap bersifat negatif. Kontribusi skor untuk butir ganjil dihitung dari jawaban dikurangi 1, sedangkan butir genap dihitung dari 5 dikurangi jawaban. Jumlah kontribusi kemudian dikalikan 2,5 sehingga skor berada pada rentang 0 sampai 100.",
        italic_terms=["System Usability Scale"],
    )
    add_rich_paragraph(
        doc,
        "Skor SUS bukan persentase jawaban benar, melainkan skor pembanding kegunaan. Untuk interpretasi deskriptif, skor sekitar 85,5 atau lebih dapat dikategorikan sangat baik; 71,4 sampai kurang dari 85,5 baik; 50,9 sampai kurang dari 71,4 cukup; dan kurang dari 50,9 buruk. Kategori tersebut mengacu pada pemetaan kata sifat yang dikembangkan dari data SUS dalam jumlah besar (Bangor, Kortum, & Miller, 2009). Usability sendiri dipahami sebagai hasil penggunaan yang mencakup efektivitas, efisiensi, dan kepuasan dalam konteks tertentu (ISO 9241-11, 2018).",
        italic_terms=["Usability"],
    )

    add_section_heading(doc, "2.10.4", "Pengujian Efisiensi Waktu", level=2)
    add_rich_paragraph(
        doc,
        "Pengujian efisiensi waktu membandingkan waktu penyelesaian tugas sebelum dan setelah menggunakan sistem. Tugas yang dibandingkan harus setara, misalnya mencatat booking, mencari ketersediaan barber, mencatat pelanggan walk-in, mengonfirmasi pembayaran, dan menyusun informasi transaksi. Tingkat penghematan waktu dihitung dengan rumus: (waktu proses lama dikurangi waktu proses sistem) dibagi waktu proses lama, kemudian dikalikan 100 persen.",
        italic_terms=["booking", "barber", "walk-in"],
    )

    add_section_heading(doc, "2.11", "Penelitian Terdahulu")
    add_rich_paragraph(
        doc,
        "Penelitian terdahulu digunakan untuk membandingkan ruang lingkup, teknologi, dan hasil yang telah dicapai. Ringkasan penelitian yang paling berkaitan dengan sistem Homcuts disajikan pada Tabel 2.1.",
    )
    add_table(
        doc,
        "Tabel 2.1 Perbandingan penelitian terdahulu",
        ["Peneliti", "Fokus dan hasil", "Keterbatasan terhadap penelitian ini"],
        [
            ["Firmansyah dkk. (2023)", "Aplikasi booking barbershop berbasis web yang mempermudah pemesanan jadwal pelanggan.", "Berfokus pada booking; belum menggabungkan POS walk-in, penjualan produk, dan antrean terpadu."],
            ["Ginoga dkk. (2023)", "Sistem informasi akuntansi barbershop untuk membantu pencatatan transaksi dan penyajian informasi keuangan.", "Berfokus pada pencatatan akuntansi; belum menempatkan jadwal layanan publik dan pelanggan dalam satu alur."],
            ["Ngatini & Cahyanti (2024)", "POS berbasis web pada Shortcut Barbershop untuk mengoptimalkan transaksi dan layanan kasir.", "Berfokus pada titik penjualan; belum menyatukan slot booking pelanggan dengan jadwal walk-in."],
            ["Apriandi dkk. (2025)", "Penerapan Laravel pada aplikasi pemesanan barbershop berbasis web untuk mengelola pemesanan dan jadwal.", "Belum menekankan integrasi antrean, POS layanan, pesanan produk tunai, dan pemutakhiran status pelanggan."],
        ],
        widths=[3.1, 6.0, 5.0],
        font_size=8.8,
    )

    add_section_heading(doc, "2.12", "Posisi Penelitian")
    add_rich_paragraph(
        doc,
        "Berdasarkan perbandingan tersebut, penelitian ini menempatkan sistem Homcuts sebagai sistem informasi manajemen barbershop terintegrasi. Kontribusi sistem terletak pada penyatuan booking pelanggan dan transaksi walk-in dalam jadwal barber serta nomor antrean yang sama; pengelolaan transaksi layanan dan produk; konfirmasi pembayaran tunai; pengelolaan data utama beserta foto; serta penyajian status kepada pelanggan dan admin melalui pembaruan berkala tanpa muat ulang manual. Dengan ruang lingkup tersebut, sistem tidak hanya menjadi media informasi atau aplikasi booking, tetapi juga mendukung operasi kasir dan pencatatan transaksi.",
        italic_terms=["barbershop", "booking", "walk-in", "barber"],
    )

    # BAB III
    doc.add_page_break()
    add_chapter_heading(doc, "BAB III", "METODE PENELITIAN")
    add_rich_paragraph(
        doc,
        "Bab ini menjelaskan tempat dan waktu penelitian, alat dan bahan, prosedur rancang bangun, penerapan rancangan ke dalam aplikasi, langkah pengujian, teknik analisis hasil pengujian, dan definisi operasional. Tahapan disusun berdasarkan pekerjaan yang benar-benar dilakukan dalam pengembangan sistem Homcuts.",
    )

    add_section_heading(doc, "3.1", "Tempat dan Waktu Penelitian")
    add_rich_paragraph(
        doc,
        "Penelitian dilaksanakan pada Homcuts Barbershop yang berlokasi di Jalan Ir. Sutami, Bulurokeng, Kota Makassar, Sulawesi Selatan. Lokasi tersebut dipilih karena menjadi tempat berlangsungnya proses pelayanan yang diteliti, meliputi penerimaan pelanggan, penentuan jadwal barber, layanan pelanggan booking dan walk-in, penjualan produk, pembayaran tunai, serta pencatatan transaksi.",
        italic_terms=["Barbershop", "barber", "booking", "walk-in"],
    )
    add_rich_paragraph(
        doc,
        "Kegiatan penelitian dan pengembangan dilaksanakan pada tahun 2026. Tahapan kegiatan meliputi identifikasi kebutuhan, pemodelan dan perancangan, implementasi aplikasi, pengujian, serta evaluasi dan perbaikan.",
    )

    add_section_heading(doc, "3.2", "Alat dan Bahan")
    add_section_heading(doc, "3.2.1", "Alat", level=2)
    add_rich_paragraph(doc, "Alat yang digunakan untuk mengembangkan dan menguji sistem disajikan pada Tabel 3.1.")
    add_table(
        doc,
        "Tabel 3.1 Alat penelitian",
        ["No.", "Alat", "Versi/jenis", "Kegunaan"],
        [
            ["1", "Komputer/laptop", "Perangkat pengembangan", "Menulis program, menjalankan container, dan melakukan pengujian."],
            ["2", "PHP", "8.4 pada container aplikasi", "Menjalankan aplikasi Laravel pada sisi server."],
            ["3", "Laravel", "13.26.1", "Kerangka kerja aplikasi web dengan pola MVC."],
            ["4", "Composer", "2", "Mengelola dependensi PHP."],
            ["5", "Node.js", "22", "Menjalankan proses pembangunan aset."],
            ["6", "Vite dan Tailwind CSS", "Vite 8; Tailwind CSS 4", "Membangun aset CSS dan JavaScript serta menyusun antarmuka."],
            ["7", "MySQL", "8.4", "Menyimpan data utama dan data transaksi."],
            ["8", "Docker dan Docker Compose", "Engine dan plugin Compose", "Menjalankan layanan web, basis data, dan scheduler secara terisolasi."],
            ["9", "PHPUnit", "12.5", "Menjalankan pengujian otomatis pada aplikasi Laravel."],
            ["10", "Peramban web", "Peramban modern", "Mengakses dan menguji antarmuka pelanggan dan admin."],
        ],
        widths=[0.8, 3.3, 3.1, 6.9],
        font_size=8.5,
    )

    add_section_heading(doc, "3.2.2", "Bahan", level=2)
    add_rich_paragraph(doc, "Bahan penelitian berupa data dan aturan operasional yang digunakan dalam sistem, sebagaimana dirangkum pada Tabel 3.2.")
    add_table(
        doc,
        "Tabel 3.2 Bahan penelitian",
        ["No.", "Bahan", "Isi"],
        [
            ["1", "Data profil usaha", "Nama usaha, deskripsi, alamat, kontak, media sosial, dan jam operasional."],
            ["2", "Data barber", "Nama, spesialisasi, foto, biodata, status aktif, serta jam mulai dan selesai kerja."],
            ["3", "Data layanan", "Nama layanan yang unik, harga, durasi, deskripsi, dan status ketersediaan."],
            ["4", "Data produk", "Nama, kategori, foto, harga, stok, ukuran, deskripsi, dan status tampil."],
            ["5", "Data galeri", "Foto hasil potongan, nama model rambut, deskripsi, kredit, dan barber terkait."],
            ["6", "Aturan operasional", "Jam buka 07.00–22.00, ketersediaan berdasarkan durasi dan barber, antrean terpadu, serta pembayaran tunai."],
            ["7", "Data transaksi", "Booking, transaksi POS, pesanan produk, item, jumlah, status pelayanan, dan status pembayaran."],
        ],
        widths=[0.8, 3.5, 9.8],
        font_size=8.8,
    )

    add_section_heading(doc, "3.3", "Prosedur Penelitian")
    add_rich_paragraph(
        doc,
        "Penelitian ini merupakan penelitian rancang bangun perangkat lunak dengan alur iteratif. Tahapan dilaksanakan secara terarah untuk menghasilkan keluaran yang dapat diperiksa, sedangkan hasil evaluasi dapat mengembalikan pekerjaan ke tahap perancangan atau implementasi. Pola tersebut memungkinkan perbaikan berulang berdasarkan temuan pengujian dan masukan pengguna.",
    )
    add_figure(doc, diagram["procedure"], "Gambar 3.1 Tahapan penelitian rancang bangun", width_cm=14.6)
    add_rich_paragraph(
        doc,
        "Gambar 3.1 menunjukkan lima tahap penelitian. Identifikasi kebutuhan menghasilkan daftar kebutuhan sistem. Pemodelan dan perancangan menghasilkan model proses, data, navigasi, arsitektur, dan antarmuka. Implementasi menerjemahkan rancangan ke dalam aplikasi. Pengujian memeriksa fungsi, kegunaan, dan efisiensi. Temuan pada pengujian menjadi dasar evaluasi dan perbaikan sebelum sistem dinyatakan siap digunakan.",
    )

    add_section_heading(doc, "3.3.1", "Analisis Kebutuhan", level=2)
    add_rich_paragraph(
        doc,
        "Kebutuhan sistem diperoleh melalui pengamatan terhadap alur pelayanan Homcuts, komunikasi langsung dengan pihak usaha, serta pemeriksaan data dan media yang digunakan dalam kegiatan operasional. Pengamatan difokuskan pada cara pelanggan memperoleh informasi, pencatatan jadwal, penerimaan pelanggan walk-in, pemberian nomor antrean, penjualan produk, konfirmasi pembayaran, dan pembuatan laporan transaksi.",
        italic_terms=["walk-in"],
    )
    add_rich_paragraph(
        doc,
        "Hasil identifikasi menunjukkan dua kelompok pengguna. Pelanggan dapat mengakses halaman publik tanpa akun, membuat booking, memesan produk, mengirim pesan, serta melihat status transaksi melalui tautan yang diberikan sistem. Admin/Kasir harus login untuk mengakses ringkasan, POS, booking, transaksi, dan pengelolaan data. Kebutuhan fungsional utama disajikan pada Tabel 3.3.",
        italic_terms=["login", "booking"],
    )
    add_table(
        doc,
        "Tabel 3.3 Kebutuhan fungsional sistem",
        ["Kode", "Pengguna", "Kebutuhan fungsional"],
        [
            ["F-01", "Pelanggan", "Melihat beranda, profil, layanan dan harga, barber, galeri, lokasi, serta kontak."],
            ["F-02", "Pelanggan", "Membuat booking dengan memilih layanan, barber tertentu atau barber yang tersedia, tanggal, dan waktu."],
            ["F-03", "Sistem", "Memeriksa jam operasional, jam kerja barber, durasi layanan, dan benturan jadwal."],
            ["F-04", "Sistem", "Membuat transaksi, pembayaran tunai tertunda, dan nomor antrean terpadu untuk layanan."],
            ["F-05", "Pelanggan", "Memilih produk, mengelola keranjang, checkout, dan melihat status pesanan."],
            ["F-06", "Sistem", "Memvalidasi stok, menghitung ulang total, mencatat item, dan menyesuaikan persediaan."],
            ["F-07", "Pelanggan", "Mengirim pesan melalui halaman kontak."],
            ["F-08", "Admin/Kasir", "Login dan logout secara aman."],
            ["F-09", "Admin/Kasir", "Mencatat pelanggan walk-in, layanan, barber, waktu, produk tambahan, dan diskon melalui POS."],
            ["F-10", "Admin/Kasir", "Melihat dan mengubah booking tanpa menimbulkan benturan jadwal."],
            ["F-11", "Sistem", "Memperbarui informasi pelanggan apabila jadwal booking diubah admin."],
            ["F-12", "Admin/Kasir", "Mengonfirmasi pembayaran booking melalui Booking atau Transaksi dan pembayaran produk melalui Transaksi."],
            ["F-13", "Admin/Kasir", "Melihat riwayat transaksi dan mengurutkan tabel langsung dari judul kolom."],
            ["F-14", "Admin", "Melakukan CRUD data barber, layanan, produk, galeri, pesan, dan pengaturan."],
            ["F-15", "Admin", "Mengunggah foto barber, produk, dan galeri."],
            ["F-16", "Sistem", "Memperbarui notifikasi, tabel, dan status secara berkala tanpa muat ulang manual."],
        ],
        widths=[1.2, 2.5, 10.4],
        font_size=8.3,
    )
    add_rich_paragraph(doc, "Kebutuhan nonfungsional yang digunakan sebagai batas kualitas sistem disajikan pada Tabel 3.4.")
    add_table(
        doc,
        "Tabel 3.4 Kebutuhan nonfungsional sistem",
        ["Kode", "Aspek", "Kebutuhan"],
        [
            ["NF-01", "Keamanan", "Halaman admin dilindungi autentikasi, otorisasi admin, validasi server, CSRF, dan pembatasan percobaan login."],
            ["NF-02", "Integritas", "Perubahan transaksi yang saling terkait diproses dalam transaksi basis data dan menggunakan relasi kunci asing."],
            ["NF-03", "Konsistensi jadwal", "Booking dan POS menggunakan pemeriksaan interval serta sumber antrean yang sama."],
            ["NF-04", "Kegunaan", "Antarmuka responsif, berbahasa Indonesia, dan menampilkan pesan kesalahan yang jelas."],
            ["NF-05", "Kinerja", "Pemeriksaan ketersediaan, pemuatan tabel, dan pembaruan status memberikan respons dalam waktu yang wajar pada lingkungan uji."],
            ["NF-06", "Keandalan data", "Data MySQL dan berkas unggahan disimpan pada volume agar bertahan ketika container dibuat ulang."],
            ["NF-07", "Kompatibilitas", "Sistem dapat digunakan pada peramban modern melalui komputer maupun telepon pintar."],
        ],
        widths=[1.2, 2.8, 10.1],
        font_size=8.5,
    )
    doc.add_page_break()
    add_figure(doc, diagram["use_case"], "Gambar 3.2 Use case diagram sistem Homcuts", width_cm=14.6)
    add_rich_paragraph(
        doc,
        "Gambar 3.2 memperlihatkan batas interaksi sistem. Pelanggan tidak diwajibkan memiliki akun karena kebutuhan utama pada sisi publik adalah memperoleh informasi, membuat booking, memesan produk, mengirim pesan, dan memantau status. Admin/Kasir menggunakan autentikasi untuk menjalankan fungsi operasional dan pengelolaan data. Login merupakan prasyarat akses admin, bukan tujuan akhir proses bisnis.",
        italic_terms=["booking", "Login"],
    )

    add_section_heading(doc, "3.3.2", "Pemodelan dan Perancangan Sistem", level=2)
    add_list_item(doc, "a)", "Perancangan proses")
    add_rich_paragraph(
        doc,
        "Perancangan proses memodelkan tiga alur utama, yaitu booking, transaksi layanan walk-in melalui POS, dan pemesanan produk. Ketiga alur menggunakan transaksi sebagai catatan keuangan, tetapi hanya booking dan POS layanan yang menggunakan jadwal barber serta nomor antrean.",
        italic_terms=["booking", "walk-in", "barber"],
    )
    add_figure(doc, diagram["activity_booking"], "Gambar 3.3 Activity diagram booking dan pembayaran tunai", width_cm=10.8)
    add_rich_paragraph(
        doc,
        "Pada Gambar 3.3, sistem tidak langsung menyimpan pilihan waktu sebelum ketersediaannya divalidasi. Booking yang berhasil menghasilkan transaksi, pembayaran tunai berstatus tertunda, dan nomor antrean. Jika admin mengubah jadwal, data yang ditampilkan pada halaman status pelanggan ikut berubah. Pembayaran dinyatakan lunas setelah kasir mengonfirmasi penerimaan uang.",
        italic_terms=["Booking"],
    )
    add_figure(doc, diagram["activity_pos"], "Gambar 3.4 Activity diagram transaksi walk-in melalui POS", width_cm=9.8)
    add_rich_paragraph(
        doc,
        "Gambar 3.4 menunjukkan bahwa pelanggan walk-in didaftarkan oleh kasir sebelum menerima layanan. Waktu layanan diperiksa terhadap booking dan transaksi walk-in lain pada barber yang sama. Setelah data disimpan, pelanggan memperoleh nomor antrean terpadu. Pembayaran tunai dikonfirmasi setelah layanan diterima sesuai alur operasional Homcuts.",
        italic_terms=["walk-in", "booking", "barber"],
    )
    add_figure(doc, diagram["activity_product"], "Gambar 3.5 Activity diagram pemesanan produk", width_cm=12.5)
    add_rich_paragraph(
        doc,
        "Pada Gambar 3.5, total dan persediaan diperiksa kembali pada server saat checkout. Pesanan produk tidak menggunakan pengiriman; pelanggan membayar tunai di lokasi. Setelah kasir mengonfirmasi pembayaran melalui halaman Transaksi, status pesanan berubah menjadi siap diambil dan selanjutnya dapat ditandai selesai setelah produk diserahkan.",
        italic_terms=["checkout"],
    )

    add_list_item(doc, "b)", "Perancangan basis data")
    add_rich_paragraph(
        doc,
        "Basis data dirancang dengan memisahkan data utama, data aktivitas, dan data pendukung. Data utama meliputi BARBERS, SERVICES, PRODUCTS, GALLERY_ENTRIES, dan SITE_SETTINGS. Data aktivitas meliputi BOOKINGS, ORDERS, ORDER_ITEMS, PAYMENTS, dan CONTACT_MESSAGES. USERS, NOTIFICATIONS, serta QUEUE_SEQUENCES mendukung autentikasi, pemberitahuan, dan penomoran antrean. Gambar 3.6 menampilkan entitas inti dan atribut yang relevan dengan proses transaksi; tabel pendukung tetap menjadi bagian basis data fisik meskipun tidak seluruhnya ditampilkan agar diagram terbaca pada kertas A4.",
    )
    add_figure(doc, diagram["erd"], "Gambar 3.6 Entity relationship diagram sistem Homcuts", width_cm=15.1)
    add_rich_paragraph(
        doc,
        "Gambar 3.6 menunjukkan bahwa satu booking memiliki paling banyak satu order, sedangkan satu order dapat memiliki beberapa item dan beberapa catatan upaya pembayaran. ORDER_ITEMS menyimpan snapshot nama, harga, jumlah, dan total baris agar riwayat transaksi tidak berubah ketika data utama diperbarui. Jadwal layanan disimpan pada BOOKING dan dicerminkan pada ORDER untuk menyatukan booking serta POS dalam pemeriksaan jadwal. QUEUE_SEQUENCES digunakan untuk menghasilkan nomor antrean layanan yang tidak duplikat.",
        italic_terms=["booking", "order", "snapshot"],
    )

    add_list_item(doc, "c)", "Perancangan struktur navigasi")
    add_figure(doc, diagram["navigation"], "Gambar 3.7 Struktur navigasi sistem Homcuts", width_cm=14.5)
    add_rich_paragraph(
        doc,
        "Struktur pada Gambar 3.7 memisahkan navigasi publik dan navigasi admin. Pada admin, urutan menu operasional utama adalah Ringkasan, Kasir POS, Booking, dan Transaksi; menu data utama berada setelahnya. Booking dan Transaksi diberi penekanan karena paling sering digunakan untuk pengelolaan jadwal serta konfirmasi pembayaran.",
        italic_terms=["Booking"],
    )

    add_list_item(doc, "d)", "Perancangan arsitektur sistem")
    add_figure(doc, diagram["architecture"], "Gambar 3.8 Rancangan arsitektur sistem Homcuts", width_cm=14.5)
    add_rich_paragraph(
        doc,
        "Gambar 3.8 menggambarkan aliran permintaan dari peramban menuju route dan middleware Laravel, kemudian ke controller dan layanan bisnis. Model Eloquent mengakses MySQL, sedangkan foto disimpan melalui Laravel Storage. Scheduler menjalankan pekerjaan berkala. Seluruh komponen utama dijalankan melalui Docker Compose sehingga konfigurasi aplikasi, basis data, dan proses terjadwal dapat direproduksi.",
        italic_terms=["route", "middleware", "controller", "Scheduler"],
    )

    add_list_item(doc, "e)", "Perancangan antarmuka")
    add_figure(doc, diagram["ui_customer"], "Gambar 3.9 Rancangan antarmuka pelanggan", width_cm=15.0)
    add_rich_paragraph(
        doc,
        "Rancangan pelanggan pada Gambar 3.9 menggunakan navigasi sederhana, informasi utama yang mudah dipindai, tombol tindakan yang jelas, dan formulir booking dengan petunjuk alur. Status ketersediaan ditempatkan dekat pilihan waktu agar pelanggan memperoleh umpan balik sebelum mengirim data.",
        italic_terms=["booking"],
    )
    add_figure(doc, diagram["ui_admin"], "Gambar 3.10 Rancangan antarmuka admin dan kasir", width_cm=15.0)
    add_rich_paragraph(
        doc,
        "Rancangan admin pada Gambar 3.10 menempatkan ringkasan pekerjaan tertunda pada dashboard dan menyediakan halaman POS untuk transaksi walk-in. Tabel admin dirancang tetap berada dalam lebar layar, dapat diurutkan melalui judul kolom, dan diperbarui secara berkala tanpa menampilkan formulir filter tambahan.",
        italic_terms=["dashboard", "walk-in"],
    )

    add_section_heading(doc, "3.3.3", "Implementasi Sistem", level=2)
    add_rich_paragraph(
        doc,
        "Implementasi dilakukan dengan Laravel menggunakan pola MVC. Route mendefinisikan alamat halaman publik dan admin. Middleware autentikasi dan otorisasi membatasi akses admin. Controller menangani permintaan dan validasi, sedangkan model Eloquent merepresentasikan tabel serta relasinya. Tampilan dibangun dengan Blade dan Tailwind CSS; JavaScript digunakan untuk keranjang, pemeriksaan jadwal, pengurutan tabel, notifikasi, dan pembaruan status berkala.",
        italic_terms=["Route", "Middleware", "Controller", "model", "Blade", "JavaScript"],
    )
    add_rich_paragraph(
        doc,
        "Logika yang dipakai oleh lebih dari satu alur ditempatkan pada kelas layanan. BookingAvailabilityService memeriksa jam kerja dan benturan interval; QueueNumberService menghasilkan nomor antrean layanan; BookingTransactionService menyelaraskan booking dengan transaksi; PaymentService mengelola status pembayaran tunai dan konsekuensinya pada status pesanan; sedangkan AdminNotifier membuat pemberitahuan untuk admin. Pemisahan ini mencegah aturan penting ditulis berulang pada controller.",
        italic_terms=["BookingAvailabilityService", "QueueNumberService", "BookingTransactionService", "PaymentService", "AdminNotifier", "controller"],
    )
    add_rich_paragraph(
        doc,
        "Pembuatan booking dan POS menggunakan waktu mulai serta waktu selesai berdasarkan durasi layanan. Sistem mengunci data yang relevan di dalam transaksi basis data ketika memeriksa ketersediaan dan menetapkan antrean. Pendekatan ini digunakan untuk mengurangi risiko dua permintaan bersamaan memperoleh slot, nomor antrean, atau stok yang sama. Transaksi produk menghitung ulang total dan memeriksa stok pada server sebelum data disimpan.",
        italic_terms=["booking"],
    )
    add_rich_paragraph(
        doc,
        "Pembayaran pada versi penelitian dibatasi pada tunai. Booking dan transaksi POS menghasilkan status pembayaran belum dibayar. Setelah kasir menekan tindakan konfirmasi, PaymentService mencatat waktu dan pengguna yang mengonfirmasi, mengubah status pembayaran menjadi lunas, serta menyelaraskan status booking atau pesanan. Perubahan tersebut dibaca oleh halaman pelanggan dan admin melalui permintaan berkala sehingga tampilan berubah tanpa muat ulang manual.",
        italic_terms=["Booking", "PaymentService", "booking"],
    )
    add_rich_paragraph(
        doc,
        "Foto barber, produk, dan galeri diterima melalui formulir unggah dengan validasi jenis serta ukuran berkas. Lokasi berkas disimpan pada basis data, sedangkan isi berkas disimpan pada storage aplikasi. Aplikasi dijalankan melalui Docker Compose yang terdiri atas container web, MySQL, dan scheduler. Volume digunakan untuk mempertahankan data MySQL dan media unggahan.",
        italic_terms=["barber", "storage", "container"],
    )

    add_section_heading(doc, "3.3.4", "Evaluasi dan Perbaikan", level=2)
    add_rich_paragraph(
        doc,
        "Evaluasi dilakukan dengan membandingkan implementasi terhadap kebutuhan fungsional, hasil pengujian, dan alur operasional Homcuts. Temuan yang menunjukkan ketidaksesuaian diperbaiki pada rancangan atau kode, kemudian diuji kembali. Contoh perbaikan yang termasuk dalam tahap ini adalah penyatuan nomor antrean booking dan POS, pemeriksaan benturan berdasarkan durasi, penyesuaian urutan navigasi admin, konfirmasi pembayaran produk pada halaman Transaksi, serta pembaruan data tanpa muat ulang manual.",
        italic_terms=["booking"],
    )

    add_section_heading(doc, "3.4", "Langkah Pengujian Sistem")
    add_section_heading(doc, "3.4.1", "Pengujian Black-Box", level=2)
    add_rich_paragraph(
        doc,
        "Pengujian black-box dilaksanakan dengan menyiapkan kondisi awal, memasukkan data, menjalankan fungsi, dan membandingkan keluaran aktual dengan keluaran yang diharapkan. Setiap skenario dinyatakan lulus apabila seluruh hasil yang diamati sesuai dengan kebutuhan. Rencana kasus uji inti disajikan pada Tabel 3.5.",
        italic_terms=["black-box"],
    )
    add_table(
        doc,
        "Tabel 3.5 Rencana pengujian black-box",
        ["Kode", "Skenario", "Hasil yang diharapkan"],
        [
            ["B-01", "Mengakses seluruh halaman publik", "Halaman tampil dan mengambil data aktif dari basis data."],
            ["B-02", "Membuat booking pada slot tersedia", "Booking, transaksi, pembayaran tunai tertunda, notifikasi, dan nomor antrean terbentuk."],
            ["B-03", "Memilih waktu di luar jam barber atau membuat interval bertabrakan", "Permintaan ditolak dan pesan ketersediaan ditampilkan."],
            ["B-04", "Membuat dua layanan bersebelahan", "Slot kedua diterima bila waktu mulai sama dengan waktu selesai slot pertama."],
            ["B-05", "Mencatat POS pada jadwal yang bertabrakan dengan booking", "Transaksi ditolak dan data lama tidak berubah."],
            ["B-06", "Membuat booking dan POS secara berurutan", "Nomor antrean layanan berbeda dan berasal dari urutan yang sama."],
            ["B-07", "Mengubah jadwal booking melalui admin", "Jadwal baru divalidasi dan halaman status pelanggan menampilkan perubahan."],
            ["B-08", "Checkout produk dengan stok cukup", "Order dan item tersimpan, total dihitung server, stok berkurang, pembayaran menunggu."],
            ["B-09", "Checkout melebihi stok", "Pesanan ditolak tanpa mengubah stok."],
            ["B-10", "Konfirmasi tunai untuk booking", "Pembayaran menjadi lunas dan booking berubah menjadi dikonfirmasi."],
            ["B-11", "Konfirmasi tunai untuk produk", "Pembayaran menjadi lunas dan pesanan berubah menjadi siap diambil."],
            ["B-12", "Konfirmasi pembayaran yang sama dua kali", "Konfirmasi kedua tidak menggandakan penerimaan atau mengubah total."],
            ["B-13", "Akses halaman admin tanpa login", "Pengguna diarahkan ke halaman login."],
            ["B-14", "Login dengan data salah dan benar", "Data salah ditolak; data benar membuka Ringkasan."],
            ["B-15", "CRUD dan unggah foto data utama", "Data tervalidasi, tersimpan, diperbarui, ditampilkan, dan dapat dinonaktifkan/dihapus sesuai aturan."],
            ["B-16", "Mengurutkan tabel dari judul kolom", "Urutan berubah langsung dan klik berikutnya membalik arah urutan."],
            ["B-17", "Membuka dua tampilan status lalu mengubah data", "Tabel, notifikasi, dan status memperoleh data terbaru tanpa muat ulang manual."],
        ],
        widths=[1.1, 5.3, 8.0],
        font_size=8.1,
    )

    add_section_heading(doc, "3.4.2", "Pengujian User Acceptance Testing", level=2)
    add_rich_paragraph(
        doc,
        "UAT dilaksanakan setelah skenario black-box utama lulus. Peserta dipilih secara purposive agar mewakili pengguna sistem, yaitu pemilik atau Admin/Kasir Homcuts dan pelanggan. Sebelum pengujian, peneliti menjelaskan tujuan serta menyediakan data uji tanpa mengarahkan langkah penyelesaian. Peserta kemudian menjalankan tugas sesuai perannya dan memberi keputusan diterima atau tidak diterima pada setiap skenario, disertai catatan apabila hasil belum sesuai dengan alur kerja yang diharapkan.",
        italic_terms=["UAT", "black-box"],
    )
    add_rich_paragraph(
        doc,
        "Sistem dinyatakan diterima apabila seluruh skenario kritis memperoleh keputusan diterima dan persentase penerimaan keseluruhan sekurang-kurangnya 80 persen. Skenario yang tidak diterima harus diperbaiki dan diuji ulang oleh pengguna. Keputusan akhir dicatat pada lembar persetujuan UAT yang memuat identitas peran peserta, tanggal pengujian, catatan, dan tanda tangan pihak Homcuts. Rencana skenario UAT disajikan pada Tabel 3.6.",
        italic_terms=["UAT"],
    )
    add_table(
        doc,
        "Tabel 3.6 Rencana pengujian User Acceptance Testing",
        ["Kode", "Aktor", "Tugas penerimaan", "Kriteria penerimaan"],
        [
            ["UAT-01", "Pelanggan", "Mencari informasi layanan, harga, barber, galeri, dan kontak.", "Informasi yang dibutuhkan dapat ditemukan dan sesuai data Homcuts."],
            ["UAT-02", "Pelanggan", "Membuat booking pada jadwal yang tersedia.", "Kode transaksi, jadwal, barber, harga, dan nomor antrean tampil dengan benar."],
            ["UAT-03", "Pelanggan", "Mencoba jadwal yang bertabrakan atau di luar jam kerja.", "Sistem menolak pilihan dan memberikan pesan yang dapat dipahami."],
            ["UAT-04", "Pelanggan", "Memesan produk melalui keranjang dan checkout.", "Rincian pesanan, total, dan petunjuk pembayaran tunai tampil benar."],
            ["UAT-05", "Pelanggan", "Memantau booking setelah admin mengubah jadwal atau pembayaran.", "Status pelanggan berubah tanpa harus memuat ulang halaman secara manual."],
            ["UAT-06", "Admin/Kasir", "Login dan memeriksa pekerjaan tertunda pada Ringkasan.", "Akses admin berhasil dan booking/pesanan tertunda mudah ditemukan."],
            ["UAT-07", "Admin/Kasir", "Mencatat pelanggan walk-in melalui POS.", "Jadwal tervalidasi dan nomor antrean tidak sama dengan transaksi layanan lain."],
            ["UAT-08", "Admin/Kasir", "Mengonfirmasi pembayaran tunai booking dan produk.", "Pembayaran menjadi lunas dan status pelayanan berubah sesuai jenis transaksi."],
            ["UAT-09", "Admin/Kasir", "Mengubah jadwal booking pelanggan.", "Benturan ditolak; perubahan yang sah tersimpan dan tampil pada pelanggan."],
            ["UAT-10", "Admin/Kasir", "Mengelola data utama, unggah foto, mengurutkan tabel, dan melihat laporan.", "Perubahan tersimpan, tabel merespons, dan riwayat transaksi dapat diperiksa."],
        ],
        widths=[1.3, 2.1, 5.2, 5.8],
        font_size=7.6,
    )

    add_section_heading(doc, "3.4.3", "Pengujian System Usability Scale", level=2)
    add_rich_paragraph(
        doc,
        "Pengujian SUS dilakukan setelah responden menjalankan tugas yang mewakili perannya. Pelanggan diminta mencari informasi, membuat booking, memesan produk, dan membaca status transaksi. Admin/Kasir diminta mencatat transaksi POS, mengelola booking, mengonfirmasi pembayaran, mengurutkan tabel, dan memperbarui data utama. Setelah tugas selesai, responden mengisi sepuluh pernyataan SUS menggunakan skala 1 (sangat tidak setuju) sampai 5 (sangat setuju).",
        italic_terms=["booking"],
    )
    add_rich_paragraph(
        doc,
        "Langkah perhitungan dilakukan untuk setiap responden: (1) mengurangi 1 dari jawaban butir ganjil; (2) mengurangi jawaban butir genap dari 5; (3) menjumlahkan seluruh kontribusi; dan (4) mengalikan jumlah tersebut dengan 2,5. Skor sistem diperoleh dari rata-rata skor seluruh responden. Identitas responden dapat disajikan secara ringkas pada Bab IV tanpa menampilkan data pribadi yang tidak diperlukan.",
    )

    add_section_heading(doc, "3.4.4", "Pengujian Efisiensi Waktu", level=2)
    add_rich_paragraph(
        doc,
        "Pengujian efisiensi dilakukan dengan mengukur durasi tugas yang sama pada proses sebelumnya dan pada sistem. Pengukuran dimulai ketika pengguna mulai memasukkan atau mencari data dan berakhir ketika hasil yang dibutuhkan tersedia. Setiap tugas dapat diulang untuk mengurangi pengaruh kesalahan pencatatan waktu. Tugas pembanding disajikan pada Tabel 3.7.",
    )
    add_table(
        doc,
        "Tabel 3.7 Tugas pengujian efisiensi waktu",
        ["Kode", "Tugas", "Titik selesai pengukuran"],
        [
            ["E-01", "Mencatat booking pelanggan", "Booking dan nomor antrean berhasil tersimpan."],
            ["E-02", "Memeriksa ketersediaan barber", "Informasi slot tersedia atau bentrok tampil."],
            ["E-03", "Mencatat pelanggan walk-in", "Transaksi POS dan jadwal layanan berhasil tersimpan."],
            ["E-04", "Mengonfirmasi pembayaran tunai", "Status pembayaran dan status pelayanan berubah."],
            ["E-05", "Mencari dan mengurutkan riwayat transaksi", "Data yang dibutuhkan tampil pada urutan yang sesuai."],
        ],
        widths=[1.2, 6.1, 7.1],
        font_size=8.7,
    )

    add_section_heading(doc, "3.5", "Teknik Analisis Hasil Pengujian")
    add_rich_paragraph(
        doc,
        "Analisis pada penelitian ini bersifat deskriptif dan hanya digunakan untuk menafsirkan hasil pengujian sistem; tidak digunakan analisis statistik inferensial. Hasil black-box diringkas sebagai jumlah kasus lulus dan gagal serta persentase keberhasilan. Kasus yang gagal dijelaskan penyebab dan perbaikannya, kemudian diuji ulang.",
        italic_terms=["black-box"],
    )
    add_rich_paragraph(
        doc,
        "Hasil UAT diringkas sebagai jumlah skenario diterima dan tidak diterima serta persentase penerimaan. Persentase dihitung dengan membagi jumlah skenario yang diterima dengan seluruh skenario yang dijalankan, kemudian dikalikan 100 persen. Catatan pengguna dianalisis untuk menentukan perbaikan; sistem belum dinyatakan diterima apabila terdapat skenario kritis yang tidak diterima.",
        italic_terms=["UAT"],
    )
    add_rich_paragraph(
        doc,
        "Skor SUS dihitung untuk setiap responden dan dirata-ratakan. Nilai rata-rata kemudian dipetakan ke kategori pada Subbab 2.10.3. Hasil efisiensi dianalisis menggunakan selisih waktu dan persentase penghematan untuk setiap tugas. Nilai positif menunjukkan sistem membutuhkan waktu lebih singkat, nilai nol menunjukkan waktu yang sama, sedangkan nilai negatif menunjukkan proses sistem lebih lambat.",
    )

    add_section_heading(doc, "3.6", "Definisi Operasional")
    add_rich_paragraph(doc, "Definisi operasional digunakan agar istilah yang diukur dan diamati memiliki arti yang konsisten. Definisi tersebut disajikan pada Tabel 3.8.")
    add_table(
        doc,
        "Tabel 3.8 Definisi operasional",
        ["Istilah", "Definisi operasional", "Indikator"],
        [
            ["Booking", "Pemesanan layanan yang dibuat sebelum kedatangan melalui pelanggan atau admin.", "Data pelanggan, layanan, barber, waktu mulai/selesai, status, dan nomor antrean tersimpan."],
            ["Walk-in/POS", "Transaksi layanan pelanggan yang dicatat oleh kasir saat pelanggan datang langsung.", "Transaksi memiliki jadwal, barber, layanan, status pembayaran, dan antrean yang sama ruang lingkupnya dengan booking."],
            ["Slot tersedia", "Interval layanan yang berada dalam jam kerja barber dan tidak beririsan dengan layanan aktif lain.", "Validasi interval menghasilkan status tersedia."],
            ["Antrean terpadu", "Penomoran layanan yang menggunakan satu sumber untuk booking dan POS.", "Tidak terdapat nomor antrean duplikat pada transaksi layanan."],
            ["Pembayaran lunas", "Uang tunai telah diterima dan dikonfirmasi oleh admin/kasir.", "Payment berstatus paid, paid_at dan confirmed_by terisi, serta order berstatus paid."],
            ["Efektivitas fungsional", "Kesesuaian fungsi dengan keluaran yang diharapkan.", "Persentase kasus black-box yang lulus."],
            ["Penerimaan pengguna", "Keputusan pengguna bahwa alur sistem sesuai dengan kebutuhan operasionalnya.", "Seluruh skenario kritis diterima dan persentase penerimaan UAT sekurang-kurangnya 80 persen."],
            ["Kegunaan", "Kemudahan penggunaan yang dirasakan setelah menjalankan tugas sistem.", "Skor rata-rata SUS 0–100."],
            ["Efisiensi waktu", "Perubahan waktu penyelesaian tugas dibanding proses sebelumnya.", "Selisih waktu dan persentase penghematan."],
        ],
        widths=[3.0, 6.2, 5.2],
        font_size=8.2,
    )

    doc.add_page_break()
    add_chapter_heading(doc, "DAFTAR RUJUKAN", "BAB II DAN BAB III")
    refs = [
        "Apriandi, M. N., Irawan, A. S. Y., & Purwantoro. (2025). Implementasi framework Laravel pada aplikasi pemesanan barbershop berbasis web (studi kasus: Maiden Barberrock). Jurnal Informatika dan Teknik Elektro Terapan, 13(3S1), 96–108. https://doi.org/10.23960/jitet.v13i3S1.7520",
        "Bangor, A., Kortum, P., & Miller, J. (2009). Determining what individual SUS scores mean: Adding an adjective rating scale. Journal of Usability Studies, 4(3), 114–123. https://uxpajournal.org/wp-content/uploads/sites/7/pdf/JUS_Bangor_May2009.pdf",
        "Brooke, J. (1996). SUS: A ‘quick and dirty’ usability scale. Dalam P. W. Jordan, B. Thomas, B. A. Weerdmeester, & I. L. McClelland (Ed.), Usability Evaluation in Industry. Taylor & Francis. https://hci-studies.org/methods-and-measures/downloads/SUS_Brooke1996.pdf",
        "Docker. (2026). Docker Compose documentation. https://docs.docker.com/compose/",
        "Firmansyah, D., Purwanto, H., Wiharko, T., & Purbayanto, B. (2023). Aplikasi booking barbershop online berbasis web. Jurnal Internal, 6(2), 146–155. https://jurnal.masoemuniversity.ac.id/index.php/internal/article/download/849/596",
        "Ginoga, W. S., Hidayat, M., & Pakaya, N. (2023). Sistem informasi akuntansi barbershop. Journal of System and Information Technology, 3(1), 27–36. https://ejurnal.ung.ac.id/index.php/diffusion/article/download/12408/5963",
        "International Organization for Standardization. (2018). ISO 9241-11:2018 Ergonomics of human-system interaction—Part 11: Usability: Definitions and concepts. https://www.iso.org/standard/63500.html",
        "International Software Testing Qualifications Board. (2024). Certified Tester Foundation Level syllabus v4.0.1. https://istqb.org/wp-content/uploads/2024/11/ISTQB_CTFL_Syllabus_v4.0.1.pdf",
        "Laravel. (2026). Laravel 13.x documentation. https://laravel.com/docs/13.x",
        "Laudon, K. C., & Laudon, J. P. (2022). Management Information Systems: Managing the Digital Firm (17th ed.). Pearson.",
        "Ngatini, N., & Cahyanti, F. L. D. (2024). Perancangan sistem point of sales (POS) berbasis web untuk optimalisasi layanan pada Shortcut Barbershop. Jurnal Nasional Komputasi dan Teknologi Informasi, 7(6), 1707–1715. https://doi.org/10.32672/jnkti.v7i6.8260",
        "Object Management Group. (2017). OMG Unified Modeling Language (OMG UML), version 2.5.1. https://www.omg.org/spec/UML/2.5.1/PDF",
        "Oracle. (2026). MySQL 8.4 Reference Manual. https://dev.mysql.com/doc/refman/8.4/en/",
        "Sommerville, I. (2016). Software Engineering (10th ed.). Pearson.",
    ]
    for ref in refs:
        add_bibliography_entry(doc, ref)

    doc.save(OUT_PATH)
    return OUT_PATH


def main() -> None:
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True, exist_ok=True)
    diagrams = make_diagrams()
    output = build_document(diagrams)
    print(output)


if __name__ == "__main__":
    main()
