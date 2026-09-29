from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\Administrator\Documents\Rogan\budi\da-project")
SOURCE = ROOT / "output" / "modul" / "Modul_Pemahaman_Proyek_HOMCUTS.docx"
OUTPUT = ROOT / "output" / "modul" / "Modul_Alur_Lengkap_Proyek_HOMCUTS.docx"


def set_paragraph_text(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            paragraph._p.remove(run._r)
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def paragraph_by_text(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def replace_exact(doc, old, new):
    paragraph_by_text(doc, old)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == old:
            set_paragraph_text(paragraph, new)
            return


def replace_cell_exact(doc, old, new):
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text == old:
                values = new.split(" | ")
                for cell, value in zip(row.cells, values):
                    set_cell_text(cell, value)
                return
            for cell in row.cells:
                if cell.text.strip() == old:
                    set_cell_text(cell, new)
                    return
    raise ValueError(f"Table cell not found: {old}")


def add_paragraph_before(anchor, style, text=None, bold_prefix=None, body=None):
    paragraph = anchor.insert_paragraph_before(style=style)
    if text is not None:
        paragraph.add_run(text)
    elif bold_prefix is not None:
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(body or "")
    return paragraph


def add_row_after(table, anchor_row, values):
    row = table.add_row()
    anchor_row._tr.addnext(row._tr)
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)
    return row


def fit_table(table, widths_inches):
    """Write consistent table-grid and cell widths for Word's fixed layout."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, width in enumerate(widths_inches):
        table.columns[index].width = Inches(width)
    for row in table.rows:
        for index, width in enumerate(widths_inches):
            row.cells[index].width = Inches(width)
    tbl_width = table._tbl.tblPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblW")
    tbl_width.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", str(round(sum(widths_inches) * 1440)))
    tbl_width.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type", "dxa")


doc = Document(SOURCE)

# Rebase audience and inspection date to the current request and source tree.
replace_exact(
    doc,
    "Alur kerja, arsitektur, dan kode aplikasi Laravel untuk pembaca yang sudah mengenal PHP, MVC, dan MySQL",
    "Alur kerja, arsitektur, dan kode aplikasi Laravel untuk pembaca yang memahami PHP dasar, MVC, dan OOP",
)
replace_exact(
    doc,
    "Disusun berdasarkan source code dan konfigurasi proyek HOMCUTS aktual\nVersi pemeriksaan: 21 September 2026",
    "Disusun berdasarkan source code dan konfigurasi proyek HOMCUTS\nPemeriksaan source: 28 September 2026",
)
replace_exact(
    doc,
    "Karena Anda sudah memahami PHP, MVC, dan MySQL, setiap konsep baru dijelaskan dengan cara menghubungkannya ke konsep tersebut. Misalnya, route dipandang sebagai pemilih controller, model Eloquent sebagai lapisan query MySQL, Blade sebagai template PHP, dan migration sebagai riwayat perubahan struktur database.",
    "Karena Anda memahami PHP dasar, MVC, dan OOP, konsep Laravel dijelaskan dengan menghubungkannya ke class, object, method, pewarisan, dan pembagian tanggung jawab. Istilah database dan SQL yang dibutuhkan juga diterangkan saat pertama kali muncul.",
)

# Correct current frontend, schedule, payment, and migration behavior.
replacements = {
    "Vite bukan framework backend dan tidak menggantikan PHP. Ia hanya alat pengembangan/build aset frontend. Inputnya adalah resources/css/app.css dan resources/js/app.js. Saat build, Vite menghasilkan file bernama hash di public/build serta manifest. @vite(...) pada layout membaca manifest agar browser memuat file yang benar.":
    "Vite bukan framework backend dan tidak menggantikan PHP. Ia adalah alat pengembangan dan build aset frontend. vite.config.js mendaftarkan resources/css/app.css, resources/js/app.js, dan resources/js/admin-image-editor.js. Saat build, Vite menghasilkan file bernama hash serta manifest di public/build. Direktif @vite(...) pada layout membaca manifest agar browser memuat file hasil build yang benar.",
    "Service db memakai image mysql:8.4 dan nama container homcuts-db. Ia membuat database utama barbershop, user aplikasi barbershop, dan menyimpan data di volume mysql_data. Healthcheck mysqladmin memastikan MySQL siap sebelum web/scheduler mulai.":
    "Service db memakai image mysql:8.4 dan nama container homcuts-db. Ia membuat database utama barbershop, user aplikasi barbershop, dan menyimpan data di volume mysql_data. Healthcheck menjalankan mysqladmin ping ke 127.0.0.1 dengan --protocol=TCP. Compose menunggu pemeriksaan koneksi TCP ini berhasil sebelum web atau scheduler dianggap boleh mulai.",
    "app.js membaca service, tanggal, waktu, mode booking, dan capster. Berdasarkan durasi layanan, batas waktu input dihitung sebagai nilai terkecil antara 21.30 dan jam 22.00 dikurangi durasi. Ini hanya bantuan antarmuka; server tetap melakukan validasi penuh.":
    "app.js membaca service, tanggal, waktu, mode booking, dan capster. Nilai jam buka, jam tutup, serta durasi global berasal dari pengaturan situs. Waktu mulai terakhir dihitung dari jam tutup dikurangi durasi layanan. Jam kerja capster dapat membatasi pilihan lebih jauh. Nilai min/max pada form hanya membantu pengguna; server tetap melakukan validasi penuh.",
    "BookingAvailabilityService mengambil durasi layanan lalu membentuk starts_at dan ends_at. Jadwal harus mulai pada/ setelah 07.00, tidak melampaui last start 21.30, dan selesai paling lambat 22.00. Jadwal juga harus berada di dalam work_start_time dan work_end_time capster.":
    "BookingAvailabilityService membaca service_duration_minutes, store_open_time, dan store_close_time melalui ServiceSchedule. Nilai awalnya 45 menit, 07.00, dan 22.00, tetapi admin dapat mengubahnya di Pengaturan. Waktu mulai terakhir adalah jam tutup dikurangi durasi global. Jadwal juga harus berada di dalam work_start_time dan work_end_time capster.",
    "PaymentService membuat Payment cash/pending dengan UUID reference dan expiry 30 menit. Jika pembuatan payment gagal, order dibatalkan agar booking tidak menjadi orphan yang memblokir slot.":
    "PaymentService membuat Payment cash/pending dengan UUID reference dan batas waktu yang dibaca PaymentPolicy. Nilai payment_expiry_minutes di Pengaturan berlaku untuk booking dan pesanan produk; bila pengaturan belum tersedia, konfigurasi memakai nilai fallback 30 menit. Nilai dibatasi antara 5 dan 1.440 menit. Booking juga menyimpan hold_expires_at. Jika pembuatan payment gagal, order dibatalkan agar booking tidak menjadi orphan yang memblokir slot.",
    "7. Buat Payment cash pending dengan tenggat 120 menit dan kirim notifikasi admin.":
    "7. Buat Payment cash pending dengan batas waktu dari PaymentPolicy dan kirim notifikasi admin. Batas waktu yang sama diatur melalui Pengaturan situs dan dipakai untuk booking maupun pesanan produk.",
    "Pelanggan membawa kode transaksi dan membayar tunai. Admin mengonfirmasi dari halaman Transaksi. PaymentService mengubah payment menjadi paid, order payment_status menjadi paid, dan order status menjadi ready. Ketika barang telah diambil, admin dapat mengubah status proses menjadi completed.":
    "Pelanggan membawa kode transaksi dan membayar tunai. Admin mengonfirmasi dari halaman Transaksi. PaymentService mengubah payment.status menjadi paid dan order.payment_status menjadi paid. Status proses order biasanya tetap pending, sehingga antarmuka menampilkan bahwa pesanan sudah dibayar tetapi masih menunggu diambil. Setelah barang diserahkan, admin menandai status proses completed.",
    "Migration dijalankan untuk mengubah struktur database, bukan setiap kali request. Tabel migrations mencatat migration yang sudah “Ran”. Saat container web mulai, artisan migrate --force hanya menjalankan migration baru. Semua migration proyek saat pemeriksaan sudah berstatus Ran.":
    "Migration mengubah struktur atau data awal database. Tabel migrations mencatat migration yang sudah berhasil dijalankan. Saat container web mulai, artisan migrate --force memeriksa daftar tersebut dan menjalankan migration yang masih pending. Status aktual dapat dilihat dengan php artisan migrate:status; jangan menganggap semua environment memiliki status yang sama.",
}
for old, new in replacements.items():
    replace_exact(doc, old, new)

replace_exact(
    doc,
    "Istilah baru akan lebih mudah dipahami jika dipetakan ke konsep PHP, MVC, dan MySQL yang sudah Anda kenal.",
    "Istilah Laravel baru akan lebih mudah dipahami jika dihubungkan ke PHP dasar, MVC, dan OOP yang sudah Anda kenal.",
)

replace_exact(
    doc,
    "Perintah default container test menjalankan tests/Feature/BlackBoxTest.php. Untuk seluruh suite, command dapat dioverride menjadi php vendor/bin/phpunit --do-not-cache-result. Verifikasi terakhir menghasilkan 67 test dan 774 assertion lulus; black-box menghasilkan 16 test dan 104 assertion lulus.",
    "Perintah default container test menjalankan tests/Feature/BlackBoxTest.php. Untuk seluruh suite, command container dapat dioverride menjadi php vendor/bin/phpunit --do-not-cache-result. Service test memakai database barbershop_testing; jangan arahkan RefreshDatabase ke database utama.",
)
replace_exact(doc, "14.4 Hasil verifikasi saat modul dibuat", "14.4 Menjalankan pengujian")
replace_exact(
    doc,
    "Menjalankan pengujian dari folder da-project",
    "Contoh perintah pengujian dari root proyek",
)

# Give the reader a complete trace through the request/response boundary.
blade_heading = paragraph_by_text(doc, "2.3 Apa sebenarnya Blade?")
add_paragraph_before(blade_heading, "Heading 3", "Urutan lengkap dari request sampai response")
add_paragraph_before(
    blade_heading,
    "Normal",
    "Bayangkan pengguna memilih jadwal booking lalu menekan tombol Kirim. Browser tidak memanggil class Laravel secara langsung. Browser mengirim HTTP request ke web server; Laravel kemudian menentukan route, menjalankan kode PHP, dan mengirim HTTP response. Setiap langkah di bawah adalah bagian dari satu perjalanan itu.",
)

steps = [
    ("Browser membuat request. ", "Request berisi method seperti GET atau POST, path seperti /bookings, query string, header, cookie session, dan bila form dikirim, body berisi nilai input serta token CSRF. Pemeriksaan slot dari JavaScript memakai GET dan Accept: application/json; pengiriman booking memakai POST."),
    ("Docker meneruskan koneksi. ", "Pada penggunaan lokal, port host 8000 diteruskan ke port 8000 pada container web. Server PHP di container menerima koneksi. Browser tidak berbicara langsung dengan MySQL; Laravel yang mengakses service db melalui jaringan Compose."),
    ("Front controller menerima request. ", "Semua request web masuk melalui public/index.php. File ini memeriksa mode maintenance, memuat Composer autoloader, membuat aplikasi dari bootstrap/app.php, menangkap request sebagai object Illuminate\\Http\\Request, lalu menyerahkannya ke aplikasi untuk ditangani."),
    ("Bootstrap menyiapkan aplikasi. ", "bootstrap/app.php mendaftarkan routes/web.php, routes/console.php, alias middleware admin, dan aturan exception. Provider mendaftarkan layanan aplikasi. Route definitions menjadi peta yang dapat dicocokkan router dengan URL dan method HTTP."),
    ("Router memilih route. ", "Router mencocokkan method dan path request. GET /gallery menuju PageController::gallery; POST /bookings menuju BookingController::store. Nama route seperti bookings.store membantu membuat URL tanpa menulis path berulang kali."),
    ("Middleware menyaring request. ", "Route web memakai middleware untuk cookie, session, dan CSRF. Route tertentu menambahkan throttle; route admin menambahkan auth dan admin. Middleware dapat meneruskan request, mengalihkan pengguna ke login, atau menghentikan request dengan status seperti 403, 419, atau 429."),
    ("Container membangun controller. ", "Laravel membuat controller dan memasukkan dependency yang diminta di constructor. BookingController menerima Request, BookingAvailabilityService, BookingTransactionService, PaymentService, dan AdminNotifier. Untuk /payments/{payment}, route model binding mencari Payment memakai reference karena model Payment menentukan getRouteKeyName()."),
    ("Controller memvalidasi input. ", "Controller mengambil nilai dari Request dan memeriksa aturan seperti required, format tanggal, nilai yang diperbolehkan, serta keberadaan service aktif. Jika validasi gagal pada form web, Laravel mengembalikan pengguna ke halaman sebelumnya dengan pesan error dan old input di session. Jika request mengharapkan JSON, error dapat dikirim sebagai JSON dengan status 422."),
    ("Aturan bisnis dijalankan. ", "Controller memanggil service untuk aturan yang lebih besar. BookingAvailabilityService menghitung durasi, jam buka, jam kerja capster, dan konflik jadwal. Untuk menyimpan booking, DB::transaction membungkus perubahan; lockForUpdate melindungi data penting dari dua request bersamaan."),
    ("Eloquent berbicara dengan MySQL. ", "Model seperti Booking, Order, OrderItem, Product, dan Payment mewakili tabel. Query Eloquent diterjemahkan menjadi SQL. Perubahan dalam transaksi disimpan bersama-sama; bila terjadi exception sebelum transaksi selesai, MySQL membatalkan perubahan dalam transaksi tersebut."),
    ("Controller memilih jenis response. ", "Controller dapat mengembalikan view, redirect, JSON, atau file PDF. Laravel mengubah hasil controller menjadi HTTP response lengkap dengan status code, header, cookie, session, dan body. Exception yang tidak ditangani diproses oleh exception handler."),
    ("Blade membentuk HTML. ", "Untuk view, Blade menerima data controller dan layout. View composer di AppServiceProvider membagikan pengaturan situs serta katalog produk kepada view. Sintaks {{ ... }} meng-escape teks. Blade dirender di server; browser menerima HTML biasa, bukan file Blade."),
    ("Browser memuat aset lalu melanjutkan interaksi. ", "Direktif @vite membaca manifest public/build dan memasukkan alamat CSS/JavaScript hasil build ke HTML. Browser meminta aset tersebut secara terpisah. Setelah halaman tampil, app.js dapat mengirim fetch baru untuk cek slot atau status; setiap fetch memulai request HTTP baru dari awal alur ini."),
]
for prefix, body in steps:
    add_paragraph_before(blade_heading, "List Number", bold_prefix=prefix, body=body)

add_paragraph_before(blade_heading, "Heading 3", "Bentuk response yang dikirim kembali")
response_items = [
    ("HTML halaman. ", "PageController::home() mengembalikan view('home', data); hasil akhirnya biasanya status 200 dengan HTML."),
    ("Redirect. ", "BookingController::store() mengembalikan to_route('payments.show', $payment) dan flash message. Browser menerima redirect, lalu mengirim GET baru ke halaman payment."),
    ("JSON. ", "BookingController::availability() dan PaymentController::status() mengembalikan response()->json(...). JavaScript membaca body JSON tanpa memuat ulang seluruh halaman."),
    ("PDF. ", "PaymentController::invoice() merender view invoice melalui Dompdf, lalu mengirim body PDF dengan Content-Type application/pdf dan nama file unduhan."),
]
for prefix, body in response_items:
    add_paragraph_before(blade_heading, "List Bullet", bold_prefix=prefix, body=body)

# Add the now-current Compose startup behavior and the root cause of the recent incident.
chapter4 = next(p for p in doc.paragraphs if p.text.strip() == "BAB 4")
add_paragraph_before(chapter4, "Heading 2", "3.6 Urutan startup web dan kesiapan MySQL")
startup_paragraphs = [
    "Compose memulai db lebih dahulu dan menunggu status service_healthy sebelum menjalankan web. Setelah MySQL siap, web membuat storage link, menjalankan php artisan migrate --force, membersihkan compiled view, lalu menjalankan php artisan serve. Scheduler menunggu web sehat karena membutuhkan aplikasi dan database yang sama.",
    "Healthcheck saat ini memaksa koneksi TCP: mysqladmin ping -h 127.0.0.1 --protocol=TCP. Ini penting saat volume database masih kosong. Image MySQL menyalakan server sementara untuk inisialisasi; pemeriksaan lama ke localhost dapat memakai Unix socket dan berhasil sebelum MySQL membuka port TCP 3306. Compose lalu menganggap db sehat terlalu awal. Pemeriksaan TCP menunggu server normal benar-benar menerima koneksi.",
    "Perintah Artisan juga mem-bootstrap aplikasi Laravel. Saat route didaftarkan, routes/web.php meminta AdminResources::keys(); definisi resource membaca pengaturan lewat ServiceSchedule::settings() dan Schema::hasTable('site_settings'). Jika MySQL belum menerima koneksi, web dapat gagal sebelum migration dimulai. Dengan healthcheck TCP, bootstrap terjadi setelah koneksi database tersedia.",
    "Error Duplicate column name image_path berbeda dari Connection refused. Error duplikat muncul bila struktur dalam volume MySQL sudah memiliki kolom yang hendak ditambahkan migration yang masih dianggap pending. docker compose up --build membangun ulang image, tetapi mempertahankan volume mysql_data. Karena itu build ulang bukan reset database. Jangan menghapus volume kecuali datanya memang boleh hilang.",
]
for paragraph_text in startup_paragraphs:
    add_paragraph_before(chapter4, "Normal", paragraph_text)

# Update remaining business rules that changed in the September migrations.
replace_cell_exact(
    doc,
    "Layanan | Nama unik, slug, durasi, harga, deskripsi, urutan, aktif. | Daftar layanan, booking, POS.",
    "Layanan | Nama unik, slug, harga, deskripsi, urutan, aktif. Durasi operasional diatur global. | Daftar layanan, booking, POS.",
)
replace_cell_exact(
    doc,
    "services | Katalog layanan. | name/slug unique; duration; price.",
    "services | Katalog layanan. | name/slug unique; price; durasi 45 menit menjadi data legacy. Durasi aktif dibaca dari site_settings global.",
)
replace_cell_exact(
    doc,
    "site_settings | Konfigurasi teks situs. | key unique, value, group.",
    "site_settings | Konfigurasi situs dan operasional. | key unique, value, group; termasuk jam buka/tutup, durasi layanan global, dan batas pembayaran.",
)
replace_cell_exact(
    doc,
    "Pengaturan | Key/value informasi umum, kontak, sosial, dan jam. | Footer, kontak, identitas situs.",
    "Pengaturan | Key/value informasi umum, kontak, sosial, jam buka/tutup, durasi layanan global, dan batas pembayaran. | Footer, kontak, jadwal booking, dan tenggat transaksi.",
)

# Replace stale test totals with repeatable instructions rather than unverified results.
test_table = doc.tables[37]
for cell, value in zip(test_table.rows[0].cells, ["Suite", "Cakupan", "Cara menjalankan"]):
    set_cell_text(cell, value)
for row, values in zip(test_table.rows[1:], [
    ["Black-box", "Request/response publik dan admin, database, validasi, booking, stok, pembayaran.", "docker compose --profile test run --rm test"],
    ["Seluruh suite", "Feature test dan unit test.", "Override command service test ke php vendor/bin/phpunit --do-not-cache-result; gunakan barbershop_testing."],
]):
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)

# Fix and extend the endpoint index to match routes/web.php.
routes = doc.tables[42]
route_edits = {
    "/admin/{resource}/create": ["GET", "/admin/{resource}/create", "Admin\\ResourceController@create", "Tampilkan form resource"],
}
for row in routes.rows:
    cells = [c.text.strip() for c in row.cells]
    if cells[1] == "/admin/{resource}/create":
        for cell, value in zip(row.cells, route_edits[cells[1]]):
            set_cell_text(cell, value)
    elif cells[1] == "/admin/{resource}/{record}/edit":
        for cell, value in zip(row.cells, ["GET", "/admin/{resource}/{record}/edit", "Admin\\ResourceController@edit", "Tampilkan form edit"]):
            set_cell_text(cell, value)
    elif cells[1] == "/admin/{resource}/{record}":
        if cells[0] == "GET/PUT":
            for cell, value in zip(row.cells, ["PUT", "/admin/{resource}/{record}", "Admin\\ResourceController@update", "Simpan perubahan"]):
                set_cell_text(cell, value)
    elif cells[1] == "/admin/{resource}":
        if cells[0] == "GET":
            for cell, value in zip(row.cells, ["GET", "/admin/{resource}", "Admin\\ResourceController@index", "Daftar resource"]):
                set_cell_text(cell, value)

last_route_row = routes.rows[-1]
new_routes = [
    ["POST", "/admin/{resource}", "Admin\\ResourceController@store", "Simpan resource baru"],
    ["GET", "/payments/{payment}/invoice.pdf", "PaymentController@invoice", "Unduh invoice PDF"],
    ["POST", "/admin/orders/{order}/confirm-deposit", "Admin\\PaymentController@confirmDeposit", "Konfirmasi DP booking"],
    ["POST", "/admin/orders/{order}/complete", "Admin\\PaymentController@complete", "Selesaikan transaksi/layanan"],
]
for route_values in new_routes:
    last_route_row = add_row_after(routes, last_route_row, route_values)

# Bring the source inventory forward to the live settings and migrations.
inventory = doc.tables[45]
admin_resource_row = next(row for row in inventory.rows if row.cells[0].text.strip() == "`app/Support/AdminResources.php`")
support_rows = [
    ["`app/Support/PaymentPolicy.php`", "Aturan batas waktu payment dan perhitungan DP."],
    ["`app/Support/ServiceSchedule.php`", "Mengambil durasi global, jam buka/tutup, dan waktu mulai terakhir."],
]
anchor = admin_resource_row
for values in support_rows:
    anchor = add_row_after(inventory, anchor, values)

migration_anchor = next(row for row in inventory.rows if row.cells[0].text.strip() == "`database/migrations/2026_09_15_001100_rename_visible_barber_terms_to_capster.php`")
migration_rows = [
    ["`database/migrations/2026_09_22_001300_standardize_service_duration_and_order_status.php`", "Menyeragamkan data durasi layanan dan status order."],
    ["`database/migrations/2026_09_22_001400_add_global_service_duration_setting.php`", "Menambahkan pengaturan durasi layanan global."],
    ["`database/migrations/2026_09_22_001500_add_store_operating_hours_settings.php`", "Menambahkan pengaturan jam buka dan tutup."],
    ["`database/migrations/2026_09_22_001600_add_payment_expiry_setting.php`", "Menambahkan pengaturan batas waktu pembayaran."],
]
anchor = migration_anchor
for values in migration_rows:
    anchor = add_row_after(inventory, anchor, values)

# Add terminology that helps a PHP/OOP reader map Laravel's conventions.
glossary = doc.tables[49]
anchor = glossary.rows[-1]
for values in [
    ["Dependency injection", "Laravel memasukkan object yang dibutuhkan ke constructor atau method, alih-alih class membuat semua dependency sendiri."],
    ["Service container", "Object registry Laravel yang mengetahui cara membuat controller, service, dan dependency lain."],
    ["Flash message", "Data session sementara, misalnya pesan sukses yang ditampilkan setelah redirect."],
    ["HTTP response", "Object hasil yang membawa status, header, cookie, dan body kembali ke browser."],
]:
    anchor = add_row_after(glossary, anchor, values)

doc.core_properties.title = "Modul Alur Lengkap Proyek HOMCUTS"
doc.core_properties.subject = "Laravel request-response, MVC, booking, transaksi, frontend, dan Docker"
doc.core_properties.keywords = "HOMCUTS, Laravel, PHP, MVC, OOP, Blade, Vite, Tailwind, request, response"

# These appendix tables had inconsistent Word table-grid and cell widths.
# Give Word explicit, matching widths so the content stays inside page margins.
for table_index, widths in {
    30: [1.55, 2.00, 2.65],        # active/generated/legacy file categories
    31: [3.10, 3.10],               # internal naming note
    42: [0.70, 1.55, 2.05, 1.90],  # route catalogue
    43: [3.10, 3.10],               # route note callout
    47: [2.05, 1.00, 1.40, 1.70],  # class/method index
    48: [3.10, 3.10],               # index note callout
    49: [1.55, 4.65],               # glossary
    51: [3.10, 3.10],               # closing callout
}.items():
    fit_table(doc.tables[table_index], widths)

# Keep the two long appendix callouts and the command block together across pages.
for table_index in [48, 50]:
    for row in doc.tables[table_index].rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
paragraph_by_text(doc, "D.1 Cheat sheet").paragraph_format.page_break_before = True
# Prevent Word from joining the two adjacent tables and repeating the method
# index header above the explanatory callout on its continuation page.
doc.tables[48]._tbl.addprevious(OxmlElement("w:p"))

doc.save(OUTPUT)

# The source product workflow diagram contained an obsolete fixed 120-minute
# payment timeout. Update its label to match PaymentPolicy/settings.
with ZipFile(SOURCE, "r") as source_package:
    product_flow = Image.open(BytesIO(source_package.read("word/media/image6.png"))).convert("RGBA")
draw = ImageDraw.Draw(product_flow)
draw.rectangle((190, 1033, 493, 1100), fill=(255, 245, 217, 255))
font_path = Path(r"C:\Windows\Fonts\arial.ttf")
font = ImageFont.truetype(str(font_path), 24)
label_lines = ["Pembayaran tunai pending", "Tenggat di Pengaturan"]
line_height = 32
top = 1040
for line in label_lines:
    box = draw.textbbox((0, 0), line, font=font)
    text_width = box[2] - box[0]
    x = (product_flow.width - text_width) / 2
    draw.text((x, top), line, font=font, fill=(16, 16, 16, 255))
    top += line_height
diagram_bytes = BytesIO()
product_flow.save(diagram_bytes, format="PNG")

temporary_package = OUTPUT.with_suffix(".patched.docx")
with ZipFile(OUTPUT, "r") as package_in, ZipFile(temporary_package, "w", compression=ZIP_DEFLATED) as package_out:
    for item in package_in.infolist():
        payload = diagram_bytes.getvalue() if item.filename == "word/media/image6.png" else package_in.read(item.filename)
        package_out.writestr(item, payload)
temporary_package.replace(OUTPUT)
print(f"Saved {OUTPUT}")
